from __future__ import annotations

import json
import socket
from pathlib import Path
from zipfile import ZipFile

import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from nanobot.agent.tools.context import RequestContext, bind_request_context, reset_request_context
from nanobot.agent.tools.presentation import CreatePresentationTool
from nanobot.agent.tools.presentation_export import ExportPresentationTool, _preview_input
from nanobot.presentation_kimi import export_local
from nanobot.presentations import (
    KIMI_ASSETS,
    PresentationError,
    PresentationService,
    presentation_runtime_lines,
)
from nanobot.security.protection import SecurityPolicyStore


def write_sources(folder, *, chart_type="bar", horizontal=False, stack=None):
    (folder / "pages").mkdir(exist_ok=True)
    (folder / "media").mkdir(exist_ok=True)
    Image.new("RGB", (200, 100), "#E1E5EA").save(folder / "media/photo.png")
    manifest = {
        "version": "v2", "title": "本地样稿", "size": [960, 540],
        "theme": {"colors": {"accent": "#00ACEE", "primary": "#00295F"},
                  "textStyles": {"title": {"fontSize": 36, "color": "$primary", "fontFamily": "Arial"}},
                  "tableStyles": {"data": {"cellStyle": {"fontSize": 16, "border": {"color": "#D5D6D9"}},
                                            "firstRowStyle": {"fill": "$primary", "color": "#FFFFFF", "bold": True}}}},
        "pages": ["pages/cover.page", "pages/data.page"],
    }
    cover = {"notes": "示例数据，仅用于验证本地渲染", "elements": [
        {"elementId": "band", "elementType": "shape", "shapeName": "rect", "bounds": [0, 0, 960, 10],
         "fill": {"type": "gradient", "angle": 0, "stops": [{"position": 0, "color": "$accent"}, {"position": 1, "color": "$primary"}]}},
        {"elementId": "title", "elementType": "text", "bounds": [48, 60, 864, 100], "opacity": 0.9,
         "content": {"style": "$title", "text": '<p>年度<strong>经营</strong><span style="color:#00ACEE">分析</span></p><p>2026</p>'}},
        {"elementId": "photo", "elementType": "image", "bounds": [48, 200, 400, 260], "src": "media/photo.png", "fit": {"mode": "cover"}},
        {"elementId": "line", "elementType": "line", "bounds": [500, 200, 350, 1], "viewBox": [350, 1], "points": "0,0.5 350,0.5",
         "border": {"color": "$accent", "style": "dash", "width": 2}, "arrow": [None, "arrow"]},
    ]}
    series = {"type": chart_type, "name": "Revenue", "fill": "$accent", "encode": {"x": "quarter", "y": "revenue"}}
    if horizontal:
        series["encode"] = {"y": "quarter", "x": "revenue"}
    if chart_type in {"pie", "doughnut"}:
        series.update(type="pie", encode={"category": "quarter", "value": "revenue"}, fill=["$accent", "$primary", "#F17E00"])
        if chart_type == "doughnut":
            series["innerRadius"] = 0.6
    if stack:
        series["stack"] = stack
    chart = {"elementId": "growth", "elementType": "chart", "bounds": [48, 50, 520, 300],
             "data": {"cols": ["quarter", "revenue"], "rows": [["Q1", 12], ["Q2", 18], ["Q3", 22]]},
             "series": [series], "legend": False, "dataLabels": {"show": True, "content": "value"}}
    if horizontal:
        chart["yAxis"] = {"type": "category", "gridLine": False}
        chart["xAxis"] = {"type": "value", "min": 0, "title": "Revenue"}
    data = {"elements": [chart, {"elementId": "table", "elementType": "table", "bounds": [610, 80, 300, 160], "style": "$data",
                                "columnWidths": [0.6, 0.4], "rowHeights": [0.5, 0.5],
                                "rows": [[{"text": "Metric"}, {"text": "2026"}], [{"text": "Revenue"}, {"text": "22"}]]}]}
    for path, payload in (("deck.pptd", manifest), ("pages/cover.page", cover), ("pages/data.page", data)):
        (folder / path).write_text(yaml.safe_dump(payload, allow_unicode=True), encoding="utf-8")
    return manifest, cover, data


def deny_network(*args, **kwargs):
    raise AssertionError("Local Kimi export must not access the network")


def test_builtin_kimi_available_with_blocked_network_and_no_browser(tmp_path, monkeypatch):
    from nanobot import presentations
    monkeypatch.delenv("NANOBOT_PRESENTATION_KIMI_DIR", raising=False)
    monkeypatch.setattr(presentations, "runtime_executable", lambda name: None)
    monkeypatch.setattr(socket.socket, "connect", deny_network)
    SecurityPolicyStore(tmp_path).save({"network_block_all": True, "network_deny_domains": ["kimi.com"]})
    service = PresentationService(tmp_path)
    assert service.source("kimi") == KIMI_ASSETS.resolve()
    monkeypatch.setattr(presentations.PresentationService, "previews", lambda self, template: [])
    templates = [template for template in service.catalog()["templates"] if template["family"] == "kimi"]
    assert len(templates) == 4
    assert all(template["available"] and not template["requires_network"] for template in templates)
    for index, template in enumerate(templates):
        document = service.bind({"template_id": template["id"], "document_id": f"kimi-local-{index}"},
                                session_key="websocket:chat-a", project_root=tmp_path, title="本地样稿")
        source = Path(document["project_path"]) / ".source"
        assert (source / "LOCAL_EXPORT.md").is_file()
        assert (source / "LICENSE").is_file()
        assert not (source / "scripts/export_pptx.py").exists()
        assert not (source / "SKILL.md").exists()
        assert "LOCAL_EXPORT.md" in "\n".join(presentation_runtime_lines({"presentation": document}))
        write_sources(source.parent)
        assert export_local(source.parent / "deck.pptd", source.parent / "presentation.pptx") == 2


@pytest.mark.parametrize("kind,horizontal,stack,expected", [
    ("bar", False, None, XL_CHART_TYPE.COLUMN_CLUSTERED),
    ("bar", True, None, XL_CHART_TYPE.BAR_CLUSTERED),
    ("bar", False, "percent", XL_CHART_TYPE.COLUMN_STACKED_100),
    ("line", False, None, XL_CHART_TYPE.LINE),
    ("area", False, "value", XL_CHART_TYPE.AREA_STACKED),
    ("pie", False, None, XL_CHART_TYPE.PIE),
    ("doughnut", False, None, XL_CHART_TYPE.DOUGHNUT),
])
def test_native_editable_elements_export_without_network(tmp_path, monkeypatch, kind, horizontal, stack, expected):
    monkeypatch.setattr(socket.socket, "connect", deny_network)
    write_sources(tmp_path, chart_type=kind, horizontal=horizontal, stack=stack)
    output = tmp_path / "presentation.pptx"
    assert export_local(tmp_path / "deck.pptd", output) == 2
    pptx = Presentation(output)
    cover, data = pptx.slides
    assert cover.shapes[1].text == "年度经营分析\n2026"
    assert cover.shapes[1].text_frame.paragraphs[0].runs[1].font.bold
    assert cover.shapes[2].image.blob == (tmp_path / "media/photo.png").read_bytes()
    assert cover.shapes[2].crop_left == pytest.approx((1 - (400 / 260) / 2) / 2, abs=1e-5)
    assert cover.shapes[2].crop_top == 0
    assert "示例数据" in cover.notes_slide.notes_text_frame.text
    assert data.shapes[0].chart.chart_type == expected
    assert tuple(data.shapes[0].chart.series[0].values) == (12, 18, 22)
    assert data.shapes[1].table.cell(1, 1).text == "22"
    with ZipFile(output) as archive:
        assert any(name.startswith("ppt/embeddings/") for name in archive.namelist())
        assert not any(b'TargetMode="External"' in archive.read(name) for name in archive.namelist() if name.endswith(".rels"))


@pytest.mark.parametrize("path", ["https://example.com/photo.png", "../private.png", "/private.png", ".source/secret.png"])
def test_local_media_rejects_remote_and_outside_paths(tmp_path, path):
    _, cover, _ = write_sources(tmp_path)
    cover["elements"][2]["src"] = path
    (tmp_path / "pages/cover.page").write_text(yaml.safe_dump(cover))
    with pytest.raises(ValueError, match="local file|inside"):
        export_local(tmp_path / "deck.pptd", tmp_path / "presentation.pptx")
    assert not (tmp_path / "presentation.pptx").exists()


def test_local_rejects_symlinks_aliases_and_unsupported_elements(tmp_path):
    _, cover, _ = write_sources(tmp_path)
    (tmp_path / "media/link.png").symlink_to(tmp_path / "media/photo.png")
    cover["elements"][2]["src"] = "media/link.png"
    (tmp_path / "pages/cover.page").write_text(yaml.safe_dump(cover))
    with pytest.raises(ValueError, match="symlinks"):
        export_local(tmp_path / "deck.pptd", tmp_path / "presentation.pptx")
    cover["elements"][2] = {"elementId": "icon", "elementType": "icon", "bounds": [48, 200, 100, 100]}
    (tmp_path / "pages/cover.page").write_text(yaml.safe_dump(cover))
    with pytest.raises(ValueError, match="Unsupported local element: icon"):
        export_local(tmp_path / "deck.pptd", tmp_path / "presentation.pptx")
    (tmp_path / "pages/cover.page").write_text("elements: &recursive [*recursive]")
    with pytest.raises(ValueError, match="aliases"):
        export_local(tmp_path / "deck.pptd", tmp_path / "presentation.pptx")


def test_preview_copy_removes_notes_without_changing_delivered_pptx(tmp_path):
    write_sources(tmp_path)
    output = tmp_path / "presentation.pptx"
    export_local(tmp_path / "deck.pptd", output)
    before = output.read_bytes()
    preview = _preview_input(Presentation(output), tmp_path / "preview-input.pptx")
    assert output.read_bytes() == before
    assert "示例数据" in Presentation(output).slides[0].notes_slide.notes_text_frame.text
    with ZipFile(preview) as archive:
        assert not any(name.startswith(("ppt/notesSlides/", "ppt/notesMasters/")) for name in archive.namelist())
    assert len(Presentation(preview).slides) == 2


@pytest.mark.asyncio
async def test_bound_local_export_and_failed_revision_preserves_last_file(tmp_path, monkeypatch):
    from nanobot.agent.tools import presentation_export
    service = PresentationService(tmp_path)
    SecurityPolicyStore(tmp_path).save({"network_block_all": True})
    document = service.bind({"template_id": "kimi-consulting", "document_id": "kimi-document-001"},
                            session_key="websocket:chat-a", project_root=tmp_path, title="本地样稿")
    folder = Path(document["project_path"])
    _, _, data = write_sources(folder)
    original = presentation_export._script

    async def local_script(arguments, cwd, **kwargs):
        if any("render_preview.py" in value for value in arguments):
            raise PresentationError("No PDF renderer in test")
        assert arguments[1:3] == ["-m", "nanobot.presentation_kimi"]
        assert len(arguments) == 5
        return await original(arguments, cwd, **kwargs)

    monkeypatch.setattr(presentation_export, "_script", local_script)
    token = bind_request_context(RequestContext("websocket", "chat-a", session_key="websocket:chat-a", metadata={"presentation": document}))
    try:
        tool = ExportPresentationTool(workspace=tmp_path)
        result = await tool.execute(document["document_id"])
        assert isinstance(result, dict), result
        assert json.loads(result["text"])["page_count"] == 2
        assert service.document(document["document_id"])["status"] == "ready"
        assert "export_presentation" in await CreatePresentationTool(workspace=tmp_path).execute(spec_path=str(folder / "deck.pptd"))
        before = (folder / "presentation.pptx").read_bytes()
        data["elements"][0]["series"][0]["type"] = "waterfall"
        (folder / "pages/data.page").write_text(yaml.safe_dump(data))
        assert (await tool.execute(document["document_id"])).startswith("Error:")
        assert (folder / "presentation.pptx").read_bytes() == before
    finally:
        reset_request_context(token)
