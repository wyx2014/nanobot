from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

import nanobot.agent.tools.docx as docx_tool_module
from nanobot.agent.tools.docx import CreateDocxTool


@pytest.mark.asyncio
async def test_create_docx_from_markdown(tmp_path):
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (32, 20), "#3156d3").save(assets / "revenue.png")
    source = tmp_path / "report.md"
    source.write_text(
        "# 青岛啤酒研究报告\n\n"
        "这是用于验证 **DOCX** 交付的正文。\n\n"
        "---\n\n"
        "## 核心指标\n\n"
        "| 指标 | 数值 |\n| --- | --- |\n| 营收 | 100 |\n\n"
        "![收入趋势](assets/revenue.png)\n\n"
        "- 第一条观点\n- 第二条观点\n",
        encoding="utf-8",
    )
    output = tmp_path / "report.docx"

    result = await CreateDocxTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
        title="青岛啤酒研究报告",
        template="research_report",
    )

    assert isinstance(result, dict)
    assert "DOCX created successfully" in result["text"]
    output = Path(result["files"][0]["path"])
    assert output.exists()
    assert output.stat().st_size > 4_000
    assert result["files"] == [{
        "path": str(output),
        "name": output.name,
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size": output.stat().st_size,
    }]
    document = Document(str(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "青岛啤酒研究报告" in text
    assert "DOCX" in text
    assert "报告目录" not in text
    assert "研究报告 · 多角色" not in text
    assert "TPARUYI" not in document.sections[0].header.paragraphs[0].text
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "营收"
    assert len(document.inline_shapes) == 1


@pytest.mark.asyncio
async def test_create_docx_rejects_non_text_source(tmp_path):
    source = tmp_path / "report.html"
    source.write_text("<h1>Report</h1>", encoding="utf-8")

    result = await CreateDocxTool(workspace=tmp_path).execute(source_path=str(source))

    assert "Error: render_failed" in result
    assert "Markdown or text" in result


@pytest.mark.asyncio
async def test_inline_word_uses_user_template_without_source_or_companions(tmp_path):
    template = Document(docx_tool_module._DEFAULT_TEMPLATE_PATH)
    template_bytes = docx_tool_module._DEFAULT_TEMPLATE_PATH.read_bytes()
    result = await CreateDocxTool(workspace=tmp_path, allowed_dir=tmp_path).execute(
        content="# 工作报告\n\n正文内容。\n\n## 一、工作情况\n\n### （一）进度\n\n已完成。",
        output_path="report.docx",
    )
    assert isinstance(result, dict), result
    output = Path(result["files"][0]["path"])
    assert {path.name for path in tmp_path.iterdir()} == {"tmp", output.name}
    document = Document(output)
    assert document.paragraphs[0].text == "工作报告"
    assert document.paragraphs[0].style.name == "Title"
    text = "\n".join(p.text for p in document.paragraphs)
    for sample in ("中国太平", "2024年", "采用", "普通商密", "TPARUYI", "报告目录"):
        assert sample not in text
    for name, font, size, bold in (
        ("Title", "方正小标宋简体", 22, False),
        ("Body Text", "仿宋_GB2312", 16, False),
        ("Heading 1", "黑体", 16, False),
        ("Heading 2", "楷体_GB2312", 16, True),
    ):
        style = document.styles[name]
        assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == font
        assert style.font.size == Pt(size)
        assert style.font.bold is bold
        assert style.element.xml == template.styles[name].element.xml
    assert document.styles["Body Text"].paragraph_format.line_spacing == Pt(28.95)
    section = document.sections[0]
    for attribute in ("page_width", "page_height", "top_margin", "bottom_margin",
                      "left_margin", "right_margin", "gutter"):
        assert getattr(section, attribute) == getattr(template.sections[0], attribute)
    assert section.top_margin.cm == pytest.approx(3.7, abs=0.01)
    assert section.bottom_margin.cm == pytest.approx(3.5, abs=0.01)
    assert section.left_margin.cm == pytest.approx(2.6, abs=0.01)
    assert section.right_margin.cm == pytest.approx(2.6, abs=0.01)
    assert section.gutter.cm == pytest.approx(0.2, abs=0.01)
    assert section.footer._element.xml == template.sections[0].footer._element.xml
    assert "PAGE" in section.footer._element.xml
    assert not document._element.xpath('.//w:br[@w:type="page"]')
    assert docx_tool_module._DEFAULT_TEMPLATE_PATH.read_bytes() == template_bytes


@pytest.mark.asyncio
async def test_explicit_template_preserves_page_setup_styles_and_header_footer(tmp_path):
    source_template = Document()
    source_template.sections[0].top_margin = Cm(4)
    source_template.styles["Normal"].font.name = "Courier New"
    source_template.sections[0].header.paragraphs[0].text = "用户页眉"
    source_template.sections[0].footer.paragraphs[0].text = "用户页脚"
    source_template.add_paragraph("模板示例正文")
    template_path = tmp_path / "custom.docx"
    source_template.save(template_path)
    template_bytes = template_path.read_bytes()

    result = await CreateDocxTool(workspace=tmp_path).execute(
        content="# 标题\n\n正式正文。", output_path="output.docx", template_path="custom.docx",
        classification="内部", recipient="办公室：", signatory="项目组", document_date="2026年9月20日",
    )
    assert isinstance(result, dict), result
    document = Document(result["files"][0]["path"])
    assert document.sections[0].top_margin == source_template.sections[0].top_margin
    assert document.styles["Normal"].font.name == "Courier New"
    assert document.sections[0].header.paragraphs[0].text == "用户页眉"
    assert document.sections[0].footer.paragraphs[0].text == "用户页脚"
    assert [p.text for p in document.paragraphs] == [
        "内部", "标题", "办公室：", "正式正文。", "项目组", "2026年9月20日",
    ]
    assert template_path.read_bytes() == template_bytes


@pytest.mark.asyncio
@pytest.mark.parametrize("params", [
    {},
    {"content": "正文"},
    {"content": "正文", "source_path": "source.txt", "output_path": "output.docx"},
    {"content": " ", "output_path": "output.docx"},
    {"content": "正文", "output_path": "output.html"},
    {"content": "正文", "output_path": "output.docx", "template_path": "missing.docx"},
])
async def test_invalid_word_requests_do_not_create_files(tmp_path, params):
    result = await CreateDocxTool(workspace=tmp_path).execute(**params)
    assert "Error: render_failed" in result
    assert list(tmp_path.iterdir()) == []


@pytest.mark.asyncio
async def test_missing_bundled_template_does_not_fall_back_to_another_design(tmp_path, monkeypatch):
    monkeypatch.setattr(docx_tool_module, "_DEFAULT_TEMPLATE_PATH", tmp_path / "missing.docx")
    result = await CreateDocxTool(workspace=tmp_path).execute(
        content="正文", output_path="output.docx",
    )
    assert "Error: render_failed" in result
    assert not (tmp_path / "output.docx").exists()


@pytest.mark.asyncio
async def test_word_paths_respect_workspace_and_do_not_overwrite_templates(tmp_path):
    workspace = tmp_path / "project"
    workspace.mkdir()
    template_path = tmp_path / "template.docx"
    Document().save(template_path)
    template_bytes = template_path.read_bytes()
    tool = CreateDocxTool(workspace=workspace, allowed_dir=workspace)
    for params in (
        {"output_path": "../outside.docx"},
        {"output_path": "output.docx", "template_path": "../template.docx"},
    ):
        result = await tool.execute(content="正文", **params)
        assert "Error: permission_denied" in result
    assert not (tmp_path / "outside.docx").exists()
    assert list(workspace.iterdir()) == []
    result = await CreateDocxTool(workspace=tmp_path).execute(
        content="正文", output_path="template.docx", template_path="template.docx",
    )
    assert "must not overwrite" in result
    assert template_path.read_bytes() == template_bytes
