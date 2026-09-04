import zipfile
from io import BytesIO
from pathlib import Path

import httpx
import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from nanobot.agent.skills import BUILTIN_SKILLS_DIR, SkillsLoader
from nanobot.agent.tools import presentation as presentation_module
from nanobot.agent.tools.presentation import (
    CreatePresentationTool,
    ImportPresentationAssetTool,
    PresentationScriptError,
)


def _encoded_image(format_name: str = "PNG") -> bytes:
    output = BytesIO()
    Image.new("RGB", (320, 180), (12, 116, 185)).save(output, format=format_name)
    return output.getvalue()


def _mock_presentation_http(
    monkeypatch: pytest.MonkeyPatch,
    handler,
) -> None:
    transport = httpx.MockTransport(handler)
    real_async_client = httpx.AsyncClient

    class TransportAsyncClient(real_async_client):
        def __init__(self, *args, **kwargs) -> None:
            kwargs["transport"] = transport
            super().__init__(*args, **kwargs)

    monkeypatch.setattr(presentation_module.httpx, "AsyncClient", TransportAsyncClient)
    monkeypatch.setattr(
        "nanobot.agent.tools.web._validate_url_safe",
        lambda _url: (True, ""),
    )


def _write_spec(path: Path, slides: list[dict]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        yaml.safe_dump(
            {
                "version": 1,
                "title": "2026年度经营分析",
                "organization": "太平资产管理有限公司",
                "author": "战略运营中心",
                "date": "2026年8月31日",
                "footer": "内部资料",
                "slides": slides,
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    return path


@pytest.mark.asyncio
async def test_create_presentation_from_corporate_spec(tmp_path: Path) -> None:
    source = _write_spec(
        tmp_path / "annual-report" / "deck.yaml",
        [
            {"type": "cover", "title": "信息科技部（2026年度重点项目跟踪）"},
            {
                "type": "title-body",
                "title": "本期核心结论",
                "lead": "重点项目总体按计划推进",
                "bullets": ["三项重点工程完成阶段验收", "下一阶段关注数据治理质量"],
            },
            {
                "type": "metrics",
                "title": "核心经营指标",
                "metrics": [
                    {"label": "营业收入", "value": "12.6亿元", "change": "同比 +18.2%"},
                    {"label": "综合成本率", "value": "96.3%", "change": "同比下降 1.1pct"},
                ],
            },
            {
                "type": "two-column",
                "title": "进展与计划",
                "left": {"title": "本期已完成", "bullets": ["完成需求评审", "核心模块上线"]},
                "right": {"title": "下期重点", "bullets": ["完成数据迁移", "开展用户培训"]},
            },
            {"type": "closing"},
        ],
    )
    output = source.with_name("年度经营分析.pptx")

    result = await CreatePresentationTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    assert result["files"] == [
        {
            "path": str(output),
            "name": output.name,
            "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "size": output.stat().st_size,
        }
    ]
    presentation = Presentation(output)
    assert len(presentation.slides) == 5
    assert len(presentation.slide_masters) >= 2
    assert (presentation.slide_width, presentation.slide_height) == (12192000, 6858000)
    slide_text = [
        "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        for slide in presentation.slides
    ]
    assert "2026年度重点项目跟踪" in slide_text[0]
    assert "本期核心结论" in slide_text[1]
    assert "12.6亿元" in slide_text[2]
    assert "本期已完成" in slide_text[3]
    cover_text_shapes = [shape for shape in presentation.slides[0].shapes if shape.has_text_frame]
    assert cover_text_shapes[0].top + cover_text_shapes[0].height <= cover_text_shapes[1].top


@pytest.mark.asyncio
async def test_create_presentation_rejects_media_outside_deck_project(tmp_path: Path) -> None:
    outside = tmp_path / "outside.png"
    outside.write_bytes(b"not-an-image")
    source = _write_spec(
        tmp_path / "project" / "deck.yaml",
        [
            {"type": "cover", "title": "测试演示文稿"},
            {"type": "image-text", "title": "图片页", "image": "../outside.png"},
        ],
    )
    output = source.with_suffix(".pptx")

    result = await CreatePresentationTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
    )

    assert isinstance(result, str)
    assert "media path escapes the deck project directory" in result
    assert not output.exists()


@pytest.mark.asyncio
async def test_preview_failure_keeps_valid_pptx(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = _write_spec(
        tmp_path / "preview" / "deck.yaml",
        [
            {"type": "cover", "title": "预览降级测试"},
            {"type": "closing"},
        ],
    )
    output = source.with_suffix(".pptx")
    preview = source.with_name("deck-preview.pdf")
    original = presentation_module._run_skill_script

    def fail_preview(script_name: str, arguments: list[str]) -> dict:
        if script_name == "render_preview.py":
            raise PresentationScriptError("render_failed", "LibreOffice unavailable")
        return original(script_name, arguments)

    monkeypatch.setattr(presentation_module, "_run_skill_script", fail_preview)
    result = await CreatePresentationTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
        preview_path=str(preview),
    )

    assert isinstance(result, dict), result
    assert output.is_file()
    assert not preview.exists()
    assert len(result["files"]) == 1
    assert "PDF preview unavailable" in result["text"]


def test_corporate_ppt_is_builtin_and_has_template_asset(tmp_path: Path) -> None:
    loader = SkillsLoader(tmp_path, builtin_skills_dir=BUILTIN_SKILLS_DIR)
    entries = loader.list_skills(filter_unavailable=False)
    entry = next(item for item in entries if item["name"] == "corporate-ppt")

    assert entry["source"] == "builtin"
    template = BUILTIN_SKILLS_DIR / "corporate-ppt" / "assets" / "ppt-template.pptx"
    presentation = Presentation(template)
    assert len(presentation.slides) == 3
    assert len(presentation.slide_masters) == 2


@pytest.mark.asyncio
async def test_rich_visual_slides_remain_editable(tmp_path: Path) -> None:
    project = tmp_path / "rich-deck"
    media = project / "media"
    media.mkdir(parents=True)
    for index, color in enumerate(((12, 116, 185), (0, 167, 90), (210, 225, 234)), start=1):
        Image.new("RGB", (900 + index * 100, 600), color).save(media / f"visual-{index}.jpg")

    source = _write_spec(
        project / "deck.yaml",
        [
            {"type": "cover", "title": "图文丰富汇报"},
            {
                "type": "image-grid",
                "title": "三类现场证据共同支持当前判断",
                "source": "内部项目影像库",
                "images": [
                    {"path": "media/visual-1.jpg", "caption": "现场一"},
                    {"path": "media/visual-2.jpg", "caption": "现场二"},
                    {"path": "media/visual-3.jpg", "caption": "现场三"},
                ],
            },
            {
                "type": "timeline",
                "title": "项目按四个里程碑推进",
                "events": [
                    {"period": "1月", "title": "立项", "description": "明确范围"},
                    {"period": "3月", "title": "设计", "description": "完成评审"},
                    {"period": "6月", "title": "试点", "description": "小范围验证"},
                    {"period": "9月", "title": "推广", "description": "全面上线", "highlight": True},
                ],
            },
            {
                "type": "process",
                "title": "交付过程设置四道质量门禁",
                "steps": [
                    {"title": "需求", "description": "验收口径"},
                    {"title": "设计", "description": "架构评审"},
                    {"title": "验证", "description": "质量门禁"},
                    {"title": "上线", "description": "灰度与回滚", "highlight": True},
                ],
            },
            {
                "type": "table",
                "title": "方案B在扩展能力上更符合目标",
                "columns": ["维度", "方案A", "方案B"],
                "rows": [["周期", "2个月", "4个月"], ["扩展能力", "一般", "强"]],
                "highlight_rows": [2],
                "source": "项目可研报告",
            },
            {
                "type": "chart",
                "title": "调用量增长快于成本",
                "chart_type": "line",
                "categories": ["Q1", "Q2", "Q3", "Q4"],
                "series": [
                    {"name": "调用量", "values": [100, 150, 210, 280]},
                    {"name": "成本", "values": [100, 112, 126, 145]},
                ],
                "unit": "指数",
                "body": "规模效应开始显现",
                "source": "平台监控与财务台账",
            },
            {"type": "closing"},
        ],
    )
    output = project / "rich-deck.pptx"

    result = await CreatePresentationTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    presentation = Presentation(output)
    assert len(presentation.slides) == 7
    assert sum(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for shape in presentation.slides[1].shapes
    ) == 3
    table_shape = next(shape for shape in presentation.slides[4].shapes if shape.has_table)
    assert table_shape.table.cell(2, 2).text == "强"
    chart_shape = next(shape for shape in presentation.slides[5].shapes if shape.has_chart)
    assert len(chart_shape.chart.series) == 2
    with zipfile.ZipFile(output) as archive:
        assert any(
            name.startswith("ppt/embeddings/Microsoft_Excel_")
            for name in archive.namelist()
        )


@pytest.mark.asyncio
async def test_long_content_title_stays_clear_of_template_logo(tmp_path: Path) -> None:
    title = "Deep Blue证明了搜索与专用评估在明确规则环境中的力量"
    source = _write_spec(
        tmp_path / "title-safe-area" / "deck.yaml",
        [
            {"type": "cover", "title": "标题安全区测试"},
            {"type": "title-body", "title": title, "body": "正文"},
            {"type": "closing"},
        ],
    )
    output = source.with_suffix(".pptx")

    result = await CreatePresentationTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    presentation = Presentation(output)
    title_shape = next(
        shape for shape in presentation.slides[1].shapes if getattr(shape, "text", "") == title
    )
    assert title_shape.left + title_shape.width <= Inches(10.7)


@pytest.mark.asyncio
async def test_import_presentation_asset_normalizes_local_image(tmp_path: Path) -> None:
    source = tmp_path / "incoming.bmp"
    Image.new("RGBA", (320, 180), (12, 116, 185, 180)).save(source)
    output = tmp_path / "deck" / "media" / "imported.jpg"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    assert result["files"][0]["mime_type"] == "image/jpeg"
    with Image.open(output) as imported:
        assert imported.format == "JPEG"
        assert imported.size == (320, 180)


@pytest.mark.asyncio
async def test_import_presentation_asset_blocks_private_url(tmp_path: Path) -> None:
    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_url="http://127.0.0.1/internal.png",
        output_path=str(tmp_path / "deck" / "media" / "blocked.png"),
    )

    assert isinstance(result, str)
    assert "private/internal address" in result
    assert not (tmp_path / "deck" / "media" / "blocked.png").exists()


@pytest.mark.asyncio
async def test_import_presentation_asset_uses_source_page_as_referer(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    image = _encoded_image()

    def handler(request: httpx.Request) -> httpx.Response:
        if request.headers.get("referer") != "https://archive.example/item/shakey":
            return httpx.Response(403, request=request)
        return httpx.Response(
            200,
            headers={"content-type": "image/png"},
            content=image,
            request=request,
        )

    _mock_presentation_http(monkeypatch, handler)
    output = tmp_path / "deck" / "media" / "shakey.jpg"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_url="https://cdn.example/shakey.png",
        source_page_url="https://archive.example/item/shakey",
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    assert "source: https://archive.example/item/shakey" in result["text"]
    assert output.is_file()


@pytest.mark.asyncio
async def test_import_presentation_asset_resolves_page_metadata_and_redirects(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    image = _encoded_image("WEBP")
    requests: list[tuple[str, str | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append((str(request.url), request.headers.get("referer")))
        if request.url.path == "/story":
            return httpx.Response(302, headers={"location": "/story/final"}, request=request)
        if request.url.path == "/story/final":
            html = (
                '<html><head><meta property="og:image" '
                'content="/media/alphago.webp?size=large"></head></html>'
            )
            return httpx.Response(
                200,
                headers={"content-type": "text/html; charset=utf-8"},
                content=html.encode(),
                request=request,
            )
        if request.url.path == "/media/alphago.webp":
            return httpx.Response(
                200,
                headers={"content-type": "image/webp"},
                content=image,
                request=request,
            )
        return httpx.Response(404, request=request)

    _mock_presentation_http(monkeypatch, handler)
    output = tmp_path / "deck" / "media" / "alphago.png"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_page_url="https://history.example/story",
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    assert "resolved_image_url: https://history.example/media/alphago.webp?size=large" in result["text"]
    assert requests[-1] == (
        "https://history.example/media/alphago.webp?size=large",
        "https://history.example/story",
    )
    with Image.open(output) as imported:
        assert imported.format == "PNG"
        assert imported.size == (320, 180)


@pytest.mark.asyncio
async def test_import_presentation_asset_rejects_non_image_response(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            headers={"content-type": "text/html"},
            content=b"<html>not an image</html>",
            request=request,
        )

    _mock_presentation_http(monkeypatch, handler)
    output = tmp_path / "deck" / "media" / "invalid.png"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_url="https://images.example/result",
        output_path=str(output),
    )

    assert isinstance(result, str)
    assert "non_image_response" in result
    assert not output.exists()


@pytest.mark.asyncio
async def test_import_presentation_asset_reports_missing_remote_image(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(404, request=request)

    _mock_presentation_http(monkeypatch, handler)
    output = tmp_path / "deck" / "media" / "missing.png"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_url="https://images.example/guessed-name.png",
        output_path=str(output),
    )

    assert isinstance(result, str)
    assert "http_not_found" in result
    assert "instead of guessing another URL" in result
    assert not output.exists()


@pytest.mark.asyncio
async def test_import_presentation_asset_reports_unsupported_svg(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            headers={"content-type": "image/svg+xml"},
            content=b'<svg xmlns="http://www.w3.org/2000/svg"/>',
            request=request,
        )

    _mock_presentation_http(monkeypatch, handler)
    output = tmp_path / "deck" / "media" / "diagram.png"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_url="https://images.example/diagram.svg",
        output_path=str(output),
    )

    assert isinstance(result, str)
    assert "unsupported_image" in result
    assert not output.exists()


@pytest.mark.asyncio
async def test_import_presentation_asset_reports_tls_failure_without_disabling_verification(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("SSL: CERTIFICATE_VERIFY_FAILED", request=request)

    _mock_presentation_http(monkeypatch, handler)
    output = tmp_path / "deck" / "media" / "tls.png"

    result = await ImportPresentationAssetTool(workspace=tmp_path).execute(
        source_url="https://archive.example/asset.png",
        output_path=str(output),
    )

    assert isinstance(result, str)
    assert "tls_error" in result
    assert "TLS verification was not disabled" in result
    assert not output.exists()


@pytest.mark.asyncio
async def test_create_presentation_warns_when_one_image_is_reused(tmp_path: Path) -> None:
    project = tmp_path / "repeated-media"
    media = project / "media"
    media.mkdir(parents=True)
    Image.new("RGB", (800, 600), (12, 116, 185)).save(media / "generic.jpg")
    source = _write_spec(
        project / "deck.yaml",
        [
            {"type": "cover", "title": "素材重复检查"},
            *[
                {
                    "type": "image-text",
                    "title": f"证据页{index}",
                    "image": "media/generic.jpg",
                    "body": "同一图片不应代替逐页素材策略",
                }
                for index in range(1, 4)
            ],
            {"type": "closing"},
        ],
    )
    output = project / "deck.pptx"

    result = await CreatePresentationTool(workspace=tmp_path).execute(
        source_path=str(source),
        output_path=str(output),
    )

    assert isinstance(result, dict), result
    assert "is reused on 3 slides" in result["text"]
