"""Offline PPTD adapter backed by python-pptx; see presentation_assets/kimi/LOCAL_EXPORT.md."""

from __future__ import annotations

import math
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.dml.fill import FillFormat
from pptx.dml.line import LineFormat
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

_TEXT = "fontSize fontFamily color bold italic underline lineHeight lineHeightPx"
_COMMON = "elementId elementType bounds rotation"


def _number(value, minimum=-10000, maximum=10000):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Expected a number, got {value!r}")
    if not math.isfinite(value) or not minimum <= value <= maximum:
        raise ValueError(f"Number outside [{minimum}, {maximum}]: {value}")
    return float(value)


def _fields(value, allowed, context):
    if not isinstance(value, dict):
        raise ValueError(f"{context} must be an object")
    extra = set(value) - set(allowed.split())
    if extra:
        raise ValueError(f"Unsupported local {context} fields: {sorted(extra, key=str)}; see LOCAL_EXPORT.md")
    return value


class _RichText(HTMLParser):
    def __init__(self, value):
        super().__init__(convert_charrefs=True)
        self.paragraphs = [[]]
        self.stack = []
        self.feed(value)
        self.close()
        if self.stack:
            raise ValueError("Unclosed rich-text tag")
        if len(self.paragraphs) > 1 and not self.paragraphs[-1]:
            self.paragraphs.pop()

    def handle_starttag(self, tag, attrs):
        if tag == "br":
            self.handle_data("\v")
            return
        if tag not in {"p", "span", "strong", "b", "em", "i", "u"}:
            raise ValueError(f"Unsupported rich-text tag: {tag}")
        values = _fields(dict(attrs), "style", "rich-text attribute")
        if tag == "p" and self.paragraphs[-1]:
            self.paragraphs.append([])
        style = dict(self.stack[-1][1]) if self.stack else {}
        if tag in {"strong", "b"}:
            style["bold"] = True
        if tag in {"em", "i"}:
            style["italic"] = True
        if tag == "u":
            style["underline"] = True
        for declaration in (values.get("style") or "").split(";"):
            if not declaration.strip():
                continue
            key, separator, value = declaration.partition(":")
            key, value = key.strip(), value.strip()
            if not separator:
                raise ValueError("Invalid inline text style")
            if key == "color":
                style["color"] = value
            elif key == "font-size" and re.fullmatch(r"\d+(\.\d+)?(px|pt)", value):
                style["fontSize"] = float(value[:-2])
            elif key == "font-family":
                style["fontFamily"] = value.strip("\"'")
            elif key == "font-weight" and value in {"bold", "normal", "400", "700"}:
                style["bold"] = value in {"bold", "700"}
            elif key == "font-style" and value in {"normal", "italic"}:
                style["italic"] = value == "italic"
            else:
                raise ValueError(f"Unsupported inline text style: {key}: {value}")
        self.stack.append((tag, style))

    def handle_endtag(self, tag):
        if not self.stack or self.stack[-1][0] != tag:
            raise ValueError(f"Mismatched rich-text tag: {tag}")
        self.stack.pop()
        if tag == "p":
            self.paragraphs.append([])

    def handle_data(self, data):
        if not self.stack and not data.strip() and not self.paragraphs[-1]:
            return
        self.paragraphs[-1].append((data, dict(self.stack[-1][1]) if self.stack else {}))


class LocalKimiRenderer:
    def __init__(self, manifest: Path):
        # Resolve the parent once; each relative input is still checked for traversal and symlinks.
        self.root = manifest.parent.resolve()
        self.bytes_read = 0
        self.deck = _fields(self.load(manifest.name), "version title size theme pages", "deck")
        if self.deck.get("version") != "v2":
            raise ValueError("Local PPTD requires version: v2")
        self.width, self.height = [_number(x, 72, 4032) for x in self.deck["size"]]
        self.theme = _fields(self.deck.get("theme", {}), "colors textStyles tableStyles", "theme")
        self.pptx = Presentation()
        self.pptx.slide_width, self.pptx.slide_height = Pt(self.width), Pt(self.height)
        self.pptx.core_properties.title = str(self.deck.get("title", ""))[:255]

    def local_file(self, name, limit=25 * 1024 * 1024):
        if not isinstance(name, str) or not name or ":" in name or "\\" in name:
            raise ValueError("Use a project-relative local file, not a URL")
        relative = Path(name)
        if relative.is_absolute() or any(part in {"..", ".source"} for part in relative.parts):
            raise ValueError("Media and pages must stay inside the presentation project")
        path = self.root / relative
        if path.resolve() != path or not path.resolve().is_relative_to(self.root):
            raise ValueError("Presentation files must not contain symlinks")
        size = path.stat().st_size
        self.bytes_read += size
        if size > limit or self.bytes_read > 100 * 1024 * 1024:
            raise ValueError("Presentation input exceeds the export size limit")
        return path

    def load(self, name):
        raw = self.local_file(name, 2 * 1024 * 1024).read_text(encoding="utf-8")
        # Bound YAML construction, including recursive aliases and excessive nesting.
        depth = 0
        for event in yaml.parse(raw, Loader=yaml.SafeLoader):
            if isinstance(event, yaml.AliasEvent):
                raise ValueError("YAML aliases are not supported in local PPTD")
            if isinstance(event, (yaml.MappingStartEvent, yaml.SequenceStartEvent)):
                depth += 1
                if depth > 30:
                    raise ValueError("PPTD nesting exceeds the limit")
            elif isinstance(event, (yaml.MappingEndEvent, yaml.SequenceEndEvent)):
                depth -= 1
        data = yaml.safe_load(raw)
        if not isinstance(data, dict):
            raise ValueError("PPTD file must contain an object")
        return data

    def color(self, value):
        seen = set()
        while isinstance(value, str) and value.startswith("$"):
            if value in seen:
                raise ValueError("Cyclic theme color")
            seen.add(value)
            value = self.theme.get("colors", {}).get(value[1:])
        if not isinstance(value, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}([0-9a-fA-F]{2})?", value):
            raise ValueError(f"Invalid color or theme reference: {value!r}")
        return RGBColor.from_string(value[1:7]), int(value[7:9], 16) / 255 if len(value) == 9 else 1

    def paint(self, target, value, opacity=1):
        rgb, alpha = self.color(value)
        target.rgb = rgb
        element = target._xFill.srgbClr
        for child in list(element):
            element.remove(child)
        node = OxmlElement("a:alpha")
        node.set("val", str(round(alpha * _number(opacity, 0, 1) * 100000)))
        element.append(node)

    def fill(self, target, spec, opacity=1):
        if spec is None:
            target.background()
            return
        if isinstance(spec, str):
            spec = {"type": "solid", "color": spec}
        kind = spec.get("type", "solid")
        if kind == "solid":
            _fields(spec, "type color", "solid fill")
            target.solid()
            self.paint(target.fore_color, spec["color"], opacity)
        elif kind == "gradient":
            _fields(spec, "type gradientType angle stops", "gradient")
            if spec.get("gradientType", "linear") != "linear":
                raise ValueError("Only linear gradients are supported locally")
            stops = spec["stops"]
            if not isinstance(stops, list) or not 2 <= len(stops) <= 10:
                raise ValueError("Use 2-10 gradient stops")
            target.gradient()
            target.gradient_angle = _number(spec.get("angle", 0), 0, 360)
            stop_list = target._xPr.gradFill.gsLst
            for child in list(stop_list):
                stop_list.remove(child)
            positions = []
            for stop in stops:
                _fields(stop, "position color", "gradient stop")
                position = _number(stop["position"], 0, 1)
                positions.append(position)
                rgb, alpha = self.color(stop["color"])
                node, color = OxmlElement("a:gs"), OxmlElement("a:srgbClr")
                node.set("pos", str(round(position * 100000)))
                color.set("val", str(rgb))
                opacity_node = OxmlElement("a:alpha")
                opacity_node.set("val", str(round(alpha * _number(opacity, 0, 1) * 100000)))
                color.append(opacity_node)
                node.append(color)
                stop_list.append(node)
            if positions != sorted(positions):
                raise ValueError("Gradient stops must be ordered")
        else:
            raise ValueError(f"Unsupported local fill: {kind}; use a local image element for photos")

    def border_xml(self, node, spec, opacity=1):
        if not spec:
            node.append(OxmlElement("a:noFill"))
            return
        _fields(spec, "style width color", "border")
        node.set("w", str(Pt(_number(spec.get("width", 1), 0, 100))))
        rgb, alpha = self.color(spec.get("color", "#000000"))
        fill, color, transparency = OxmlElement("a:solidFill"), OxmlElement("a:srgbClr"), OxmlElement("a:alpha")
        color.set("val", str(rgb))
        transparency.set("val", str(round(alpha * opacity * 100000)))
        color.append(transparency)
        fill.append(color)
        node.append(fill)
        styles = {"solid": "solid", "dash": "dash", "dot": "sysDot", "dashDot": "dashDot"}
        if spec.get("style", "solid") not in styles:
            raise ValueError("Unsupported border style")
        dash = OxmlElement("a:prstDash")
        dash.set("val", styles[spec.get("style", "solid")])
        node.append(dash)

    def line(self, target, spec, opacity=1):
        node = target._get_or_add_ln()
        for child in list(node):
            node.remove(child)
        self.border_xml(node, spec, opacity)

    def style(self, reference, group):
        if reference is None:
            return {}
        if not isinstance(reference, str) or not reference.startswith("$"):
            raise ValueError(f"Use a $name reference for {group}")
        styles = self.theme.get(group, {})
        if reference[1:] not in styles:
            raise ValueError(f"Unknown {group} reference: {reference}")
        return dict(styles[reference[1:]])

    def font(self, font, style, opacity=1):
        font.name = style.get("fontFamily", "Microsoft YaHei")
        font.size = Pt(_number(style.get("fontSize", 18), 1, 400))
        for key in ("bold", "italic", "underline"):
            if key in style:
                setattr(font, key, bool(style[key]))
        self.paint(font.color, style.get("color", "#111111"), opacity)
        # PowerPoint stores CJK typefaces separately from Latin typefaces.
        ea = font._rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        if ea is None:
            ea = OxmlElement("a:ea")
            font._rPr.append(ea)
        ea.set("typeface", font.name)

    def text(self, frame, content, opacity=1):
        content = {**self.style(content.get("style"), "textStyles"), **content}
        _fields(content, _TEXT + " text style align wrap", "text")
        horizontal, vertical = content.get("align", ["left", "top"])
        aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
        anchors = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        if horizontal not in aligns or vertical not in anchors:
            raise ValueError("Invalid text alignment")
        frame.clear()
        frame.word_wrap = bool(content.get("wrap", True))
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = anchors[vertical]
        value = str(content.get("text", ""))
        if len(value) > 20000:
            raise ValueError("Text element is too long")
        paragraphs = _RichText(value).paragraphs if re.search(r"</?[a-zA-Z][^>]*>", value) else [[(line, {})] for line in value.split("\n")]
        for index, runs in enumerate(paragraphs):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = aligns[horizontal]
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = (Pt(_number(content["lineHeightPx"], 1, 1000)) if "lineHeightPx" in content
                                      else _number(content.get("lineHeight", 1.15), 0.5, 5))
            for value, inline in runs:
                run = paragraph.add_run()
                run.text = value
                self.font(run.font, {**content, **inline}, opacity)

    def element(self, slide, element):
        kind = element["elementType"]
        values = element["bounds"]
        if len(values) != 4:
            raise ValueError("bounds must be [x, y, width, height]")
        x, y, width, height = [_number(value, 0, 4032) for value in values]
        if width <= 0 or height <= 0 or x + width > self.width + 0.1 or y + height > self.height + 0.1:
            raise ValueError(f"Element {element.get('elementId')} extends outside the slide")
        bounds = [Pt(value) for value in values]
        opacity = _number(element.get("opacity", 1), 0, 1)
        if kind == "text":
            _fields(element, _COMMON + " content opacity", "text element")
            shape = slide.shapes.add_textbox(*bounds)
            self.text(shape.text_frame, element["content"], opacity)
        elif kind == "shape":
            _fields(element, _COMMON + " shapeName fill border opacity adjustments", "shape")
            shape = slide.shapes.add_shape(MSO_SHAPE.from_xml(element.get("shapeName", "rect")), *bounds)
            self.fill(shape.fill, element.get("fill"), opacity)
            self.line(shape.line, element.get("border"), opacity)
            for index, adjustment in enumerate(element.get("adjustments", [])):
                shape.adjustments[index] = _number(adjustment, 0, 100000) / 100000
        elif kind == "line":
            _fields(element, _COMMON + " border opacity points viewBox arrow", "line")
            vw, vh = [_number(value, 0.001, 4032) for value in element.get("viewBox", [width, height])]
            points = element.get("points", f"0,0 {vw},{vh}").split()
            if len(points) != 2:
                raise ValueError("Use two endpoints per line; split polylines into segments")
            coordinates = [[float(value) for value in point.split(",")] for point in points]
            coordinates = [(Pt(x + _number(px, 0, vw) / vw * width), Pt(y + _number(py, 0, vh) / vh * height)) for px, py in coordinates]
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *coordinates[0], *coordinates[1])
            self.line(shape.line, element.get("border", {"color": "#111111"}), opacity)
            arrows = element.get("arrow", [None, None])
            if len(arrows) != 2:
                raise ValueError("Line arrow must have two ends")
            for tag, arrow in zip(("headEnd", "tailEnd"), arrows):
                if arrow not in {None, "arrow", "triangle", "diamond", "oval"}:
                    raise ValueError("Unsupported line arrow")
                if arrow:
                    node = OxmlElement("a:" + tag)
                    node.set("type", "triangle" if arrow == "arrow" else arrow)
                    shape.line._get_or_add_ln().append(node)
        elif kind == "image":
            _fields(element, _COMMON + " src fit opacity", "image")
            path = self.local_file(element["src"])
            fit = _fields(element.get("fit", {}), "mode", "image fit").get("mode", "cover")
            with Image.open(path) as image:
                iw, ih = image.size
                if image.format not in {"PNG", "JPEG", "GIF", "BMP", "TIFF"}:
                    raise ValueError("Use a local PNG or JPEG image")
            if fit == "contain":
                scale = min(width / iw, height / ih)
                shape = slide.shapes.add_picture(str(path), Pt(x + (width - iw * scale) / 2), Pt(y + (height - ih * scale) / 2), Pt(iw * scale), Pt(ih * scale))
            elif fit in {"cover", "fill"}:
                shape = slide.shapes.add_picture(str(path), *bounds)
                if fit == "cover":
                    if iw / ih > width / height:
                        shape.crop_left = shape.crop_right = (1 - width / height / (iw / ih)) / 2
                    else:
                        shape.crop_top = shape.crop_bottom = (1 - iw / ih / (width / height)) / 2
            else:
                raise ValueError("Image fit must be cover, contain or fill")
            if opacity != 1:
                alpha = OxmlElement("a:alphaModFix")
                alpha.set("amt", str(round(opacity * 100000)))
                shape._element.blipFill.blip.append(alpha)
        elif kind in {"table", "chart"}:
            if element.get("rotation", 0):
                raise ValueError("Native tables and charts cannot rotate")
            shape = self.table(slide, element, bounds) if kind == "table" else self.chart(slide, element, bounds)
        else:
            raise ValueError(f"Unsupported local element: {kind}; use native shapes or local images")
        shape.name = element["elementId"]
        if element.get("rotation"):
            shape.rotation = _number(element["rotation"], -360, 360)

    def table(self, slide, element, bounds):
        _fields(element, _COMMON + " rows columnWidths rowHeights style", "table")
        rows = element["rows"]
        if not isinstance(rows, list) or not 1 <= len(rows) <= 100:
            raise ValueError("Tables require 1-100 rows")
        columns = len(rows[0])
        if not 1 <= columns <= 30 or any(len(row) != columns for row in rows):
            raise ValueError("Local tables require equal-length rows and 1-30 columns")
        shape = slide.shapes.add_table(len(rows), columns, *bounds)
        table = shape.table
        table.first_row = table.horz_banding = False
        for key, tracks, count, size in (("columnWidths", table.columns, columns, bounds[2]),
                                         ("rowHeights", table.rows, len(rows), bounds[3])):
            weights = element.get(key, [1 / count] * count)
            if len(weights) != count or not math.isclose(sum(_number(x, 0.001, 1) for x in weights), 1, abs_tol=0.001):
                raise ValueError(f"{key} must be proportions summing to 1")
            sizes = [round(size * weight / sum(weights)) for weight in weights]
            sizes[-1] += size - sum(sizes)
            for track, length in zip(tracks, sizes):
                setattr(track, "width" if key == "columnWidths" else "height", length)
        theme = self.style(element.get("style"), "tableStyles")
        _fields(theme, "cellStyle firstRowStyle lastRowStyle firstColumnStyle lastColumnStyle bodyStyles rowOverColumn", "table style")
        for ri, row in enumerate(rows):
            for ci, raw in enumerate(row):
                cell = table.cell(ri, ci)
                _fields(raw, _TEXT + " text textStyle fill border align", "table cell")
                row_style, col_style = {}, {}
                if theme.get("bodyStyles") and 0 < ri < len(rows) - 1:
                    row_style.update(theme["bodyStyles"][(ri - 1) % len(theme["bodyStyles"])])
                for match, key in ((ri == 0, "firstRowStyle"), (ri == len(rows) - 1, "lastRowStyle")):
                    if match:
                        row_style.update(theme.get(key, {}))
                for match, key in ((ci == 0, "firstColumnStyle"), (ci == columns - 1, "lastColumnStyle")):
                    if match:
                        col_style.update(theme.get(key, {}))
                categories = (col_style, row_style) if theme.get("rowOverColumn", True) else (row_style, col_style)
                style = {**theme.get("cellStyle", {}), **categories[0], **categories[1],
                         **self.style(raw.get("textStyle"), "textStyles"), **raw}
                _fields(style, _TEXT + " text textStyle fill border align", "table cell style")
                self.fill(cell.fill, style.get("fill", "#FFFFFF"))
                border = style.get("border")
                sides = border if isinstance(border, list) else [border] * 4
                if len(sides) != 4:
                    raise ValueError("Cell borders use [top, right, bottom, left]")
                # CT_TableCellProperties requires left, right, top, bottom before its fill.
                for name, spec in (("lnB", sides[2]), ("lnT", sides[0]), ("lnR", sides[1]), ("lnL", sides[3])):
                    node = OxmlElement("a:" + name)
                    self.border_xml(node, spec)
                    cell._tc.get_or_add_tcPr().insert(0, node)
                content = {key: value for key, value in style.items() if key not in {"fill", "border", "textStyle"}}
                self.text(cell.text_frame, content)
                cell.margin_left = cell.margin_right = Pt(8)
                cell.margin_top = cell.margin_bottom = Pt(5)
        return shape

    def chart(self, slide, element, bounds):
        _fields(element, _COMMON + " data series legend dataLabels xAxis yAxis fontFamily title fill border", "chart")
        series = element["series"]
        if not isinstance(series, list) or not 1 <= len(series) <= 20:
            raise ValueError("Charts require 1-20 series")
        kind = series[0]["type"]
        if kind not in {"bar", "line", "area", "pie"} or any(s["type"] != kind for s in series):
            raise ValueError("Local charts support one type per chart: bar, line, area or pie")
        if kind == "pie" and len(series) != 1:
            raise ValueError("Pie charts require one series")
        if kind == "pie" and ("xAxis" in element or "yAxis" in element):
            raise ValueError("Pie charts do not have axes")
        horizontal = element.get("yAxis", {}).get("type") == "category"
        stack = series[0].get("stack")
        if stack not in {None, "value", "percent"} or any(s.get("stack") != stack for s in series):
            raise ValueError("Use the same stack mode (value or percent) for all series")
        choices = {
            "bar": (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED, XL_CHART_TYPE.BAR_STACKED_100) if horizontal else
                   (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED, XL_CHART_TYPE.COLUMN_STACKED_100),
            "line": (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_STACKED, XL_CHART_TYPE.LINE_STACKED_100),
            "area": (XL_CHART_TYPE.AREA, XL_CHART_TYPE.AREA_STACKED, XL_CHART_TYPE.AREA_STACKED_100),
            "pie": (XL_CHART_TYPE.DOUGHNUT if series[0].get("innerRadius", 0) else XL_CHART_TYPE.PIE,) * 3,
        }
        data = _fields(element["data"], "cols rows", "chart data")
        cols, rows = data["cols"], data["rows"]
        if not 1 <= len(rows) <= 500 or len(set(cols)) != len(cols) or any(len(row) != len(cols) for row in rows):
            raise ValueError("Invalid chart data dimensions or duplicate columns")
        category_key, value_key = ("category", "value") if kind == "pie" else (("y", "x") if horizontal else ("x", "y"))
        category_column = series[0]["encode"][category_key]
        chart_data = CategoryChartData()
        chart_data.categories = [str(row[cols.index(category_column)]) for row in rows]
        for item in series:
            _fields(item, "type encode name fill border innerRadius startAngle" if kind == "pie" else
                    "type encode name fill border stack", "chart series")
            _fields(item["encode"], category_key + " " + value_key, "chart encode")
            if item["encode"][category_key] != category_column:
                raise ValueError("All series must share a category column")
            values = [row[cols.index(item["encode"][value_key])] for row in rows]
            chart_data.add_series(str(item.get("name", item["encode"][value_key])),
                                  [None if value is None else _number(value, -1e15, 1e15) for value in values])
        shape = slide.shapes.add_chart(choices[kind][[None, "value", "percent"].index(stack)], *bounds, chart_data)
        chart = shape.chart
        self.font(chart.font, {"fontSize": 12, "fontFamily": element.get("fontFamily", "Microsoft YaHei")})
        # The chart-space shape properties control its frame without rasterizing chart data.
        properties = OxmlElement("c:spPr")
        chart._chartSpace.insert_element_before(properties, "c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst")
        self.fill(FillFormat.from_fill_parent(properties), element.get("fill"))
        self.line(LineFormat(properties), element.get("border"))
        legend = element.get("legend", True)
        if isinstance(legend, dict):
            _fields(legend, "show position fontSize fontFamily color bold italic", "legend")
        chart.has_legend = legend is not False and (not isinstance(legend, dict) or legend.get("show", True))
        if chart.has_legend:
            spec = legend if isinstance(legend, dict) else {}
            chart.legend.position = {"top": XL_LEGEND_POSITION.TOP, "bottom": XL_LEGEND_POSITION.BOTTOM,
                                     "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT}[spec.get("position", "bottom")]
            chart.legend.include_in_layout = False
            self.font(chart.legend.font, {"fontSize": 12, **spec})
        if element.get("title"):
            spec = element["title"]
            spec = {"text": spec} if isinstance(spec, str) else spec
            chart.has_title = True
            self.text(chart.chart_title.text_frame, {"fontSize": 18, **spec})
        palette = ["#00ACEE", "#0064BC", "#F17E00", "#D5D6D9"]
        for index, (rendered, item) in enumerate(zip(chart.series, series)):
            paint = item.get("fill", palette[index % len(palette)])
            if kind == "line":
                self.fill(rendered.format.line.fill, paint)
                rendered.format.line.width = Pt(2)
                if item.get("border"):
                    self.line(rendered.format.line, item["border"])
            elif kind == "pie":
                colors = paint if isinstance(paint, list) else (palette if "fill" not in item else [paint])
                if not colors:
                    raise ValueError("Pie colors must not be empty")
                for pi, point in enumerate(rendered.points):
                    self.fill(point.format.fill, colors[pi % len(colors)])
                    self.line(point.format.line, item.get("border"))
                if item.get("innerRadius"):
                    chart.plots[0].hole_size = round(_number(item["innerRadius"], 0.1, 0.9) * 100)
                if "startAngle" in item:
                    chart.plots[0].first_slice_angle = round(_number(item["startAngle"], 0, 360))
            else:
                self.fill(rendered.format.fill, paint)
                self.line(rendered.format.line, item.get("border"))
        labels = _fields(element.get("dataLabels", {}), "show content numberFormat fontSize fontFamily color bold italic", "data labels")
        chart.plots[0].has_data_labels = labels.get("show", False)
        if labels.get("show"):
            target = chart.plots[0].data_labels
            content = labels.get("content", "value")
            if content not in {"value", "category", "percentage"} or (content == "percentage" and kind != "pie"):
                raise ValueError("Percentage data labels are only supported on pie charts")
            target.show_value = content == "value"
            target.show_category_name = content == "category"
            target.show_percentage = content == "percentage"
            self.font(target.font, {"fontSize": 12, **labels})
            if "numberFormat" in labels:
                target.number_format = labels["numberFormat"]
        if kind != "pie":
            for key, axis, value_axis in (("yAxis" if horizontal else "xAxis", chart.category_axis, False),
                                          ("xAxis" if horizontal else "yAxis", chart.value_axis, True)):
                spec = _fields(element.get(key, {}), "type show min max title label gridLine axisLine", "axis")
                if spec.get("type", "value" if value_axis else "category") != ("value" if value_axis else "category"):
                    raise ValueError("Use one category axis and one value axis")
                axis.visible = spec.get("show", True)
                label = spec.get("label", {})
                if isinstance(label, dict):
                    _fields(label, "fontSize fontFamily color bold italic numberFormat", "axis label")
                    self.font(axis.tick_labels.font, {"fontSize": 12, **label})
                    if "numberFormat" in label:
                        axis.tick_labels.number_format = label["numberFormat"]
                elif label is False:
                    axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
                axis.has_major_gridlines = bool(spec.get("gridLine", value_axis))
                if axis.has_major_gridlines:
                    grid = spec.get("gridLine")
                    self.line(axis.major_gridlines.format.line, grid if isinstance(grid, dict) else {"color": "#E1E5EA", "width": 0.5})
                axis_line = spec.get("axisLine", True)
                self.line(axis.format.line, axis_line if isinstance(axis_line, dict) else {"color": "#9AA3AE", "width": 0.5} if axis_line else None)
                for bound in ("min", "max"):
                    if bound in spec:
                        if not value_axis:
                            raise ValueError("Axis min/max is only supported on value axes")
                        setattr(axis, "minimum_scale" if bound == "min" else "maximum_scale", _number(spec[bound], -1e15, 1e15))
                if spec.get("title"):
                    axis.has_title = True
                    title = spec["title"]
                    self.text(axis.axis_title.text_frame, {"fontSize": 12, **({"text": title} if isinstance(title, str) else title)})
        return shape

    def save(self, output: Path):
        pages = self.deck["pages"]
        if not isinstance(pages, list) or not 1 <= len(pages) <= 100 or len(set(pages)) != len(pages):
            raise ValueError("Use 1-100 unique page paths")
        for page_path in pages:
            page = _fields(self.load(page_path), "pageType background notes elements", "page")
            elements = page.get("elements", [])
            if not isinstance(elements, list) or not 1 <= len(elements) <= 500:
                raise ValueError("Each page requires 1-500 elements")
            slide = self.pptx.slides.add_slide(self.pptx.slide_layouts[6])
            self.fill(slide.background.fill, page.get("background", "#FFFFFF"))
            ids = set()
            for element in elements:
                name = element.get("elementId")
                if not isinstance(name, str) or not name or name in ids:
                    raise ValueError(f"Missing or duplicate elementId in {page_path}")
                ids.add(name)
                try:
                    self.element(slide, element)
                except (ValueError, KeyError, TypeError, AttributeError, IndexError) as exc:
                    raise ValueError(f"{page_path}, element {name}: {exc}") from exc
            if page.get("notes"):
                slide.notes_slide.notes_text_frame.text = str(page["notes"])
        self.pptx.save(output)
        return len(pages)


def export_local(manifest: Path, output: Path):
    return LocalKimiRenderer(manifest).save(output)


if __name__ == "__main__":
    export_local(Path(sys.argv[1]), Path(sys.argv[2]))
