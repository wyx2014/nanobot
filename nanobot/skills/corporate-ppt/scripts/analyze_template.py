#!/usr/bin/env python3
"""Inspect a PowerPoint template and emit a machine-readable layout catalog."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

from pptx import Presentation

from ppt_common import DEFAULT_TEMPLATE, DeckSpecError


def analyze(template_path: Path) -> dict:
    template = template_path.expanduser().resolve()
    if not template.is_file():
        raise DeckSpecError(f"PowerPoint template not found: {template}")
    presentation = Presentation(template)
    masters = []
    for master_index, master in enumerate(presentation.slide_masters):
        layouts = []
        for layout_index, layout in enumerate(master.slide_layouts):
            placeholders = []
            for shape in layout.placeholders:
                placeholder_type = str(shape.placeholder_format.type).split()[0]
                placeholders.append(
                    {
                        "index": shape.placeholder_format.idx,
                        "type": placeholder_type,
                        "name": shape.name,
                        "x": round(shape.left / 914400, 3),
                        "y": round(shape.top / 914400, 3),
                        "width": round(shape.width / 914400, 3),
                        "height": round(shape.height / 914400, 3),
                    }
                )
            layouts.append(
                {
                    "index": layout_index,
                    "name": layout.name,
                    "placeholders": placeholders,
                }
            )
        masters.append({"index": master_index, "layouts": layouts})
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        slides.append(
            {
                "index": index,
                "layout": slide.slide_layout.name,
                "text": [
                    shape.text.strip()
                    for shape in slide.shapes
                    if shape.has_text_frame and shape.text.strip()
                ],
            }
        )
    return {
        "template": str(template),
        "sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
        "slide_width_inches": round(presentation.slide_width / 914400, 3),
        "slide_height_inches": round(presentation.slide_height / 914400, 3),
        "master_count": len(presentation.slide_masters),
        "layout_count": sum(len(master.slide_layouts) for master in presentation.slide_masters),
        "slides": slides,
        "masters": masters,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Analyze a PowerPoint template.")
    parser.add_argument("template", nargs="?", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", "-o", type=Path, help="optional JSON output file")
    parser.add_argument("--force", action="store_true", help="replace an existing output file")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        result = analyze(args.template)
        rendered = json.dumps(result, ensure_ascii=False, indent=2)
        if args.output:
            output = args.output.expanduser().resolve()
            if output.exists() and not args.force:
                raise DeckSpecError(f"output already exists; pass --force to replace it: {output}")
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(rendered + "\n", encoding="utf-8")
        else:
            print(rendered)
    except (DeckSpecError, OSError, ValueError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
