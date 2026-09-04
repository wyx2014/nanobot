#!/usr/bin/env python3
"""Shared helpers for the corporate PowerPoint skill."""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any, Iterable

import yaml
from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "ppt-template.pptx"

SLIDE_TYPES = {
    "cover",
    "title-body",
    "section",
    "two-column",
    "comparison",
    "image-text",
    "image-grid",
    "metrics",
    "timeline",
    "process",
    "table",
    "chart",
    "full-image",
    "closing",
}

BRAND_BLUE = "0874B9"
BRAND_GREEN = "00A75A"
TEXT_DARK = "18324A"
TEXT_MUTED = "5D6B78"
PALE_BLUE = "EAF4FA"
PALE_GREEN = "EAF8F1"
WHITE = "FFFFFF"
FONT_NAME = "Microsoft YaHei"


class DeckSpecError(ValueError):
    """Raised when a CorporateDeck specification is invalid."""


def load_deck_spec(path: Path) -> dict[str, Any]:
    """Load and validate a JSON or YAML CorporateDeck specification."""

    source = path.expanduser().resolve()
    if not source.is_file():
        raise DeckSpecError(f"deck specification not found: {source}")
    try:
        if source.suffix.lower() == ".json":
            payload = json.loads(source.read_text(encoding="utf-8"))
        else:
            payload = yaml.safe_load(source.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, yaml.YAMLError) as exc:
        raise DeckSpecError(f"could not parse deck specification: {exc}") from exc

    if not isinstance(payload, dict):
        raise DeckSpecError("deck specification must be an object")
    if payload.get("version") != 1:
        raise DeckSpecError("deck specification version must be 1")
    slides = payload.get("slides")
    if not isinstance(slides, list) or not slides:
        raise DeckSpecError("deck specification must contain a non-empty slides list")
    if len(slides) > 100:
        raise DeckSpecError("deck specification cannot contain more than 100 slides")

    counts: dict[str, int] = {}
    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise DeckSpecError(f"slide {index} must be an object")
        slide_type = slide.get("type")
        if slide_type not in SLIDE_TYPES:
            supported = ", ".join(sorted(SLIDE_TYPES))
            raise DeckSpecError(
                f"slide {index} has unsupported type {slide_type!r}; expected one of: {supported}"
            )
        counts[slide_type] = counts.get(slide_type, 0) + 1
        if slide_type in {"cover", "closing"} and counts[slide_type] > 1:
            raise DeckSpecError(f"deck can contain at most one {slide_type} slide")
        _validate_slide(index, slide)
    return payload


def _validate_slide(index: int, slide: dict[str, Any]) -> None:
    slide_type = str(slide["type"])
    if slide_type not in {"closing", "full-image"} and not _nonempty_text(slide.get("title")):
        raise DeckSpecError(f"slide {index} ({slide_type}) requires a title")
    if slide_type in {"image-text", "full-image"} and not _nonempty_text(slide.get("image")):
        raise DeckSpecError(f"slide {index} ({slide_type}) requires an image")
    if slide_type in {"two-column", "comparison"}:
        for side in ("left", "right"):
            if not isinstance(slide.get(side), dict):
                raise DeckSpecError(f"slide {index} ({slide_type}) requires a {side} object")
    if slide_type == "metrics":
        metrics = slide.get("metrics")
        if not isinstance(metrics, list) or not metrics:
            raise DeckSpecError(f"slide {index} (metrics) requires a non-empty metrics list")
        if len(metrics) > 6:
            raise DeckSpecError(f"slide {index} (metrics) supports at most 6 metrics")
        for metric_index, metric in enumerate(metrics, start=1):
            if not isinstance(metric, dict):
                raise DeckSpecError(f"slide {index} metric {metric_index} must be an object")
            if not _nonempty_text(metric.get("label")) or not _nonempty_text(metric.get("value")):
                raise DeckSpecError(
                    f"slide {index} metric {metric_index} requires label and value"
                )
    if slide_type == "image-grid":
        images = slide.get("images")
        if not isinstance(images, list) or not 2 <= len(images) <= 6:
            raise DeckSpecError(f"slide {index} (image-grid) requires 2 to 6 images")
        for image_index, image in enumerate(images, start=1):
            if not isinstance(image, dict) or not _nonempty_text(image.get("path")):
                raise DeckSpecError(
                    f"slide {index} image {image_index} requires a path"
                )
    if slide_type == "timeline":
        _validate_named_items(index, slide, "events", 3, 8, ("period", "title"))
    if slide_type == "process":
        _validate_named_items(index, slide, "steps", 3, 6, ("title",))
    if slide_type == "table":
        columns = slide.get("columns")
        rows = slide.get("rows")
        if not isinstance(columns, list) or not 2 <= len(columns) <= 8:
            raise DeckSpecError(f"slide {index} (table) requires 2 to 8 columns")
        if not isinstance(rows, list) or not 1 <= len(rows) <= 12:
            raise DeckSpecError(f"slide {index} (table) requires 1 to 12 rows")
        if any(not isinstance(row, list) or len(row) != len(columns) for row in rows):
            raise DeckSpecError(
                f"slide {index} (table) rows must match the column count"
            )
        widths = slide.get("column_widths")
        if widths is not None and (
            not isinstance(widths, list)
            or len(widths) != len(columns)
            or any(not isinstance(value, (int, float)) or value <= 0 for value in widths)
        ):
            raise DeckSpecError(
                f"slide {index} (table) column_widths must contain one positive number per column"
            )
    if slide_type == "chart":
        chart_type = slide.get("chart_type")
        if chart_type not in {"column", "bar", "line", "pie", "doughnut"}:
            raise DeckSpecError(
                f"slide {index} (chart) chart_type must be column, bar, line, pie or doughnut"
            )
        categories = slide.get("categories")
        series = slide.get("series")
        if not isinstance(categories, list) or not 2 <= len(categories) <= 20:
            raise DeckSpecError(f"slide {index} (chart) requires 2 to 20 categories")
        if not isinstance(series, list) or not 1 <= len(series) <= 5:
            raise DeckSpecError(f"slide {index} (chart) requires 1 to 5 series")
        if chart_type in {"pie", "doughnut"} and len(series) != 1:
            raise DeckSpecError(
                f"slide {index} ({chart_type}) requires exactly one series"
            )
        for series_index, item in enumerate(series, start=1):
            values = item.get("values") if isinstance(item, dict) else None
            if (
                not isinstance(item, dict)
                or not _nonempty_text(item.get("name"))
                or not isinstance(values, list)
                or len(values) != len(categories)
                or any(not isinstance(value, (int, float)) for value in values)
            ):
                raise DeckSpecError(
                    f"slide {index} chart series {series_index} must have a name and one numeric value per category"
                )
        if not _nonempty_text(slide.get("source")):
            raise DeckSpecError(f"slide {index} (chart) requires a source")


def _validate_named_items(
    slide_index: int,
    slide: dict[str, Any],
    field: str,
    minimum: int,
    maximum: int,
    required: tuple[str, ...],
) -> None:
    items = slide.get(field)
    if not isinstance(items, list) or not minimum <= len(items) <= maximum:
        raise DeckSpecError(
            f"slide {slide_index} ({slide['type']}) requires {minimum} to {maximum} {field}"
        )
    for item_index, item in enumerate(items, start=1):
        if not isinstance(item, dict) or any(not _nonempty_text(item.get(key)) for key in required):
            joined = " and ".join(required)
            raise DeckSpecError(
                f"slide {slide_index} {field} item {item_index} requires {joined}"
            )


def _nonempty_text(value: Any) -> bool:
    return isinstance(value, str) and bool(value.strip())


def spec_warnings(spec: dict[str, Any]) -> list[str]:
    """Return content-density warnings without rejecting a valid specification."""

    warnings: list[str] = []
    image_uses: dict[str, list[int]] = {}
    content_types: list[str] = []
    visual_types = {"image-text", "image-grid", "metrics", "timeline", "process", "table", "chart", "full-image"}
    for index, slide in enumerate(spec["slides"], start=1):
        slide_type = str(slide.get("type") or "")
        if slide_type not in {"cover", "closing", "section"}:
            content_types.append(slide_type)
        title = str(slide.get("title") or "")
        if len(title) > 44:
            warnings.append(f"slide {index} title is longer than 44 characters")
        bullets = _strings(slide.get("bullets"))
        if len(bullets) > 7:
            warnings.append(f"slide {index} contains more than 7 bullets")
        if any(len(item) > 90 for item in bullets):
            warnings.append(f"slide {index} contains a bullet longer than 90 characters")
        image = slide.get("image")
        if _nonempty_text(image):
            image_uses.setdefault(str(image), []).append(index)
        for item in slide.get("images") or []:
            if isinstance(item, dict) and _nonempty_text(item.get("path")):
                image_uses.setdefault(str(item["path"]), []).append(index)
    visual_count = sum(slide_type in visual_types for slide_type in content_types)
    if len(content_types) >= 6 and visual_count / len(content_types) < 0.35:
        warnings.append(
            "fewer than 35% of content slides use a visual evidence or relationship layout"
        )
    for path, pages in sorted(image_uses.items()):
        if len(pages) > 2:
            warnings.append(
                f"media asset {path!r} is reused on {len(pages)} slides: "
                + ", ".join(str(page) for page in pages)
            )
    for start in range(max(0, len(content_types) - 2)):
        run = content_types[start : start + 3]
        if len(run) == 3 and len(set(run)) == 1:
            warnings.append(f"three consecutive content slides use the same layout: {run[0]}")
            break
    return warnings


def resolve_media_path(raw_path: str, spec_path: Path) -> Path:
    """Resolve media below the deck project directory and reject path traversal."""

    project_dir = spec_path.expanduser().resolve().parent
    candidate = Path(raw_path).expanduser()
    if candidate.is_absolute():
        raise DeckSpecError("media paths must be relative to the deck specification")
    resolved = (project_dir / candidate).resolve()
    try:
        resolved.relative_to(project_dir)
    except ValueError as exc:
        raise DeckSpecError(f"media path escapes the deck project directory: {raw_path}") from exc
    if not resolved.is_file():
        raise DeckSpecError(f"media file not found: {raw_path}")
    return resolved


def generate_presentation(
    spec: dict[str, Any],
    *,
    spec_path: Path,
    template_path: Path = DEFAULT_TEMPLATE,
) -> Presentation:
    """Build a presentation using the bundled corporate template."""

    template = template_path.expanduser().resolve()
    if not template.is_file():
        raise DeckSpecError(f"PowerPoint template not found: {template}")

    presentation = Presentation(template)
    if len(presentation.slides) < 3 or len(presentation.slide_masters) < 2:
        raise DeckSpecError(
            "corporate template must contain cover, body and closing prototypes and two masters"
        )

    cover_slide = presentation.slides[0]
    body_prototype = presentation.slides[1]
    closing_slide = presentation.slides[2]
    blank_layout = _corporate_blank_layout(presentation)
    ordered_slides = []

    content_page_number = 1
    for slide_spec in spec["slides"]:
        slide_type = slide_spec["type"]
        if slide_type == "cover":
            _render_cover(cover_slide, spec, slide_spec)
            ordered_slides.append(cover_slide)
            continue
        if slide_type == "closing":
            ordered_slides.append(closing_slide)
            continue

        slide = presentation.slides.add_slide(blank_layout)
        _render_content_slide(slide, slide_spec, spec_path)
        _add_footer(slide, spec, slide_spec, content_page_number)
        content_page_number += 1
        ordered_slides.append(slide)

    _retain_and_reorder_slides(presentation, ordered_slides)
    presentation.core_properties.title = str(spec.get("title") or "公司演示文稿")
    presentation.core_properties.subject = "Generated from the TPCowork corporate template"
    presentation.core_properties.author = str(spec.get("author") or "TPCowork")
    return presentation


def _corporate_blank_layout(presentation: Presentation):
    master = presentation.slide_masters[-1]
    candidates = []
    for layout in master.slide_layouts:
        content_placeholders = [
            shape
            for shape in layout.placeholders
            if str(shape.placeholder_format.type).split()[0]
            not in {"DATE", "FOOTER", "SLIDE_NUMBER"}
        ]
        if not content_placeholders:
            candidates.append(layout)
    if not candidates:
        raise DeckSpecError("corporate master does not contain a blank layout")
    return candidates[0]


def _retain_and_reorder_slides(presentation: Presentation, desired_slides: list[Any]) -> None:
    slide_id_list = presentation.slides._sldIdLst
    id_by_partname = {}
    for slide_id in list(slide_id_list):
        part = presentation.part.related_part(slide_id.rId)
        id_by_partname[str(part.partname)] = slide_id

    desired_ids = []
    desired_rids = set()
    for slide in desired_slides:
        slide_id = id_by_partname.get(str(slide.part.partname))
        if slide_id is None:
            raise DeckSpecError(f"could not resolve generated slide part: {slide.part.partname}")
        desired_ids.append(slide_id)
        desired_rids.add(slide_id.rId)

    removed_rids = [
        slide_id.rId for slide_id in list(slide_id_list) if slide_id.rId not in desired_rids
    ]
    for slide_id in list(slide_id_list):
        slide_id_list.remove(slide_id)
    for slide_id in desired_ids:
        slide_id_list.append(slide_id)
    for relationship_id in removed_rids:
        presentation.part.drop_rel(relationship_id)


def _render_cover(slide, deck_spec: dict[str, Any], slide_spec: dict[str, Any]) -> None:
    text_shapes = [shape for shape in slide.shapes if shape.has_text_frame]
    if len(text_shapes) < 3:
        raise DeckSpecError("corporate cover prototype must contain three text boxes")
    title = str(slide_spec["title"]).strip()
    _replace_text_preserving_format(text_shapes[0], title)
    _fit_cover_title(text_shapes[0], text_shapes[1], title)
    organization = str(
        slide_spec.get("organization")
        or deck_spec.get("organization")
        or "太平资产管理有限公司"
    ).strip()
    date = str(slide_spec.get("date") or deck_spec.get("date") or "").strip()
    _replace_text_preserving_format(text_shapes[1], organization)
    _replace_text_preserving_format(text_shapes[2], date)


def _fit_cover_title(title_shape, organization_shape, title: str) -> None:
    """Keep a wrapped cover title clear of the organization line."""

    if len(title) <= 16:
        font_size, top, height, organization_top = 38, 2.85, 0.9, 4.16
    elif len(title) <= 28:
        font_size, top, height, organization_top = 32, 2.58, 1.35, 4.22
    else:
        font_size, top, height, organization_top = 26, 2.5, 1.55, 4.28
    title_shape.top = Inches(top)
    title_shape.height = Inches(height)
    title_shape.text_frame.word_wrap = True
    title_shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    title_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    organization_shape.top = Inches(organization_top)
    first_paragraph = title_shape.text_frame.paragraphs[0]
    if first_paragraph.runs:
        run = first_paragraph.runs[0]
        run.font.name = FONT_NAME
        run.font.size = Pt(font_size)
        run.font.bold = True


def _replace_text_preserving_format(shape, text: str) -> None:
    frame = shape.text_frame
    first_paragraph = frame.paragraphs[0]
    if first_paragraph.runs:
        first_paragraph.runs[0].text = text
        for run in first_paragraph.runs[1:]:
            run.text = ""
    else:
        first_paragraph.text = text
    for paragraph in frame.paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def _render_content_slide(slide, slide_spec: dict[str, Any], spec_path: Path) -> None:
    slide_type = slide_spec["type"]
    if slide_type == "title-body":
        _render_title_body(slide, slide_spec)
    elif slide_type == "section":
        _render_section(slide, slide_spec, spec_path)
    elif slide_type == "two-column":
        _render_two_column(slide, slide_spec)
    elif slide_type == "comparison":
        _render_comparison(slide, slide_spec)
    elif slide_type == "image-text":
        _render_image_text(slide, slide_spec, spec_path)
    elif slide_type == "image-grid":
        _render_image_grid(slide, slide_spec, spec_path)
    elif slide_type == "metrics":
        _render_metrics(slide, slide_spec)
    elif slide_type == "timeline":
        _render_timeline(slide, slide_spec)
    elif slide_type == "process":
        _render_process(slide, slide_spec)
    elif slide_type == "table":
        _render_table(slide, slide_spec)
    elif slide_type == "chart":
        _render_chart(slide, slide_spec)
    elif slide_type == "full-image":
        _render_full_image(slide, slide_spec, spec_path)
    else:
        raise DeckSpecError(f"unsupported content slide type: {slide_type}")


def _render_title_body(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    lead = str(spec.get("lead") or "").strip()
    top = 1.48
    if lead:
        _add_text_box(
            slide,
            lead,
            0.95,
            top,
            11.25,
            0.65,
            size=20,
            color=BRAND_BLUE,
            bold=True,
        )
        top += 0.78
    bullets = _strings(spec.get("bullets"))
    body = str(spec.get("body") or "").strip()
    if body:
        _add_text_box(slide, body, 0.95, top, 11.2, 1.1, size=18, color=TEXT_DARK)
        top += 1.2
    if bullets:
        _add_bullet_box(slide, bullets, 1.0, top, 11.1, max(1.0, 6.15 - top), size=18)


def _render_section(slide, spec: dict[str, Any], spec_path: Path | None = None) -> None:
    title = str(spec["title"])
    subtitle = str(spec.get("subtitle") or "").strip()
    image_value = str(spec.get("image") or "").strip()
    if image_value and spec_path is not None:
        image = resolve_media_path(image_value, spec_path)
        _add_picture(
            slide,
            image,
            7.1,
            1.25,
            5.15,
            4.95,
            fit=str(spec.get("image_fit") or "cover"),
            focus_x=_ratio(spec.get("image_focus_x"), 0.5),
            focus_y=_ratio(spec.get("image_focus_y"), 0.5),
        )
        title_x, title_y, title_width, title_align = 0.95, 2.15, 5.7, PP_ALIGN.LEFT
        subtitle_x, subtitle_y, subtitle_width, subtitle_align = 0.95, 3.65, 5.7, PP_ALIGN.LEFT
    else:
        title_x, title_y, title_width, title_align = 1.15, 2.2, 11.0, PP_ALIGN.CENTER
        subtitle_x, subtitle_y, subtitle_width, subtitle_align = 1.8, 3.65, 9.7, PP_ALIGN.CENTER
    _add_text_box(
        slide,
        title,
        title_x,
        title_y,
        title_width,
        1.25,
        size=36,
        color=BRAND_BLUE,
        bold=True,
        align=title_align,
        valign=MSO_ANCHOR.MIDDLE,
    )
    if subtitle:
        _add_text_box(
            slide,
            subtitle,
            subtitle_x,
            subtitle_y,
            subtitle_width,
            0.8,
            size=19,
            color=TEXT_MUTED,
            align=subtitle_align,
            valign=MSO_ANCHOR.MIDDLE,
        )


def _render_two_column(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    left = spec["left"]
    right = spec["right"]
    _add_column(slide, left, 0.95, 1.55, 5.35, 4.8, BRAND_BLUE, PALE_BLUE)
    _add_column(slide, right, 6.75, 1.55, 5.35, 4.8, BRAND_GREEN, PALE_GREEN)


def _render_comparison(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    _add_comparison_panel(slide, spec["left"], 0.95, BRAND_BLUE, PALE_BLUE)
    _add_comparison_panel(slide, spec["right"], 6.75, BRAND_GREEN, PALE_GREEN)


def _render_image_text(slide, spec: dict[str, Any], spec_path: Path) -> None:
    _add_slide_title(slide, str(spec["title"]))
    image = resolve_media_path(str(spec["image"]), spec_path)
    image_on_left = str(spec.get("image_position") or "right").lower() == "left"
    image_x, text_x = (0.9, 7.0) if image_on_left else (6.75, 0.95)
    _add_picture(
        slide,
        image,
        image_x,
        1.55,
        5.55,
        4.55,
        fit=str(spec.get("image_fit") or "cover"),
        focus_x=_ratio(spec.get("image_focus_x"), 0.5),
        focus_y=_ratio(spec.get("image_focus_y"), 0.5),
    )
    bullets = _strings(spec.get("bullets"))
    body = str(spec.get("body") or "").strip()
    if body:
        _add_text_box(slide, body, text_x, 1.62, 5.15, 1.25, size=17, color=TEXT_DARK)
    if bullets:
        bullet_top = 3.0 if body else 1.65
        _add_bullet_box(slide, bullets, text_x, bullet_top, 5.1, 5.95 - bullet_top, size=16)
    caption = str(spec.get("caption") or "").strip()
    if caption:
        _add_text_box(
            slide,
            caption,
            image_x,
            6.12,
            5.55,
            0.35,
            size=10,
            color=TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )


def _render_image_grid(slide, spec: dict[str, Any], spec_path: Path) -> None:
    _add_slide_title(slide, str(spec["title"]))
    items = spec["images"]
    boxes = _image_grid_boxes(len(items))
    for item, (x, y, width, height) in zip(items, boxes):
        caption = str(item.get("caption") or "").strip()
        picture_height = height - (0.38 if caption else 0.0)
        image = resolve_media_path(str(item["path"]), spec_path)
        _add_picture(
            slide,
            image,
            x,
            y,
            width,
            picture_height,
            fit=str(item.get("fit") or "cover"),
            focus_x=_ratio(item.get("focus_x"), 0.5),
            focus_y=_ratio(item.get("focus_y"), 0.5),
        )
        if caption:
            _add_text_box(
                slide,
                caption,
                x,
                y + picture_height + 0.05,
                width,
                0.28,
                size=9.5,
                color=TEXT_MUTED,
                align=PP_ALIGN.CENTER,
            )


def _image_grid_boxes(count: int) -> list[tuple[float, float, float, float]]:
    left, top, width, height = 0.9, 1.45, 11.45, 4.82
    gap = 0.18
    if count == 2:
        item_width = (width - gap) / 2
        return [
            (left, top, item_width, height),
            (left + item_width + gap, top, item_width, height),
        ]
    if count == 3:
        large_width = width * 0.58
        small_width = width - large_width - gap
        small_height = (height - gap) / 2
        return [
            (left, top, large_width, height),
            (left + large_width + gap, top, small_width, small_height),
            (left + large_width + gap, top + small_height + gap, small_width, small_height),
        ]
    columns = 2 if count == 4 else 3
    rows = math.ceil(count / columns)
    item_width = (width - gap * (columns - 1)) / columns
    item_height = (height - gap * (rows - 1)) / rows
    return [
        (
            left + (index % columns) * (item_width + gap),
            top + (index // columns) * (item_height + gap),
            item_width,
            item_height,
        )
        for index in range(count)
    ]


def _render_timeline(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    lead = str(spec.get("lead") or "").strip()
    if lead:
        _add_text_box(slide, lead, 0.95, 1.25, 11.25, 0.48, size=16, color=TEXT_MUTED)
    events = spec["events"]
    axis_y = 3.52
    axis_x = 1.15
    axis_width = 10.95
    axis = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(axis_x),
        Inches(axis_y),
        Inches(axis_width),
        Inches(0.04),
    )
    axis.fill.solid()
    axis.fill.fore_color.rgb = RGBColor.from_string(BRAND_BLUE)
    axis.line.fill.background()
    spacing = axis_width / max(1, len(events) - 1)
    text_width = min(2.0, max(1.15, spacing * 0.88))
    for index, event in enumerate(events):
        x = axis_x + index * spacing
        color = BRAND_GREEN if event.get("highlight") else BRAND_BLUE
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - 0.12),
            Inches(axis_y - 0.1),
            Inches(0.24),
            Inches(0.24),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor.from_string(color)
        marker.line.color.rgb = RGBColor.from_string(WHITE)
        marker.line.width = Pt(1.5)
        above = index % 2 == 0
        period_y = 2.08 if above else 3.86
        title_y = 2.43 if above else 4.18
        description_y = 2.82 if above else 4.58
        left = max(0.82, min(12.48 - text_width, x - text_width / 2))
        _add_text_box(
            slide,
            str(event["period"]),
            left,
            period_y,
            text_width,
            0.28,
            size=12,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        _add_text_box(
            slide,
            str(event["title"]),
            left,
            title_y,
            text_width,
            0.38,
            size=14,
            color=TEXT_DARK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        description = str(event.get("description") or "").strip()
        if description:
            _add_text_box(
                slide,
                description,
                left,
                description_y,
                text_width,
                0.65,
                size=10.5,
                color=TEXT_MUTED,
                align=PP_ALIGN.CENTER,
            )


def _render_process(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    lead = str(spec.get("lead") or "").strip()
    if lead:
        _add_text_box(slide, lead, 0.95, 1.25, 11.25, 0.48, size=16, color=TEXT_MUTED)
    steps = spec["steps"]
    left, right, marker_y = 1.15, 12.05, 2.5
    spacing = (right - left) / max(1, len(steps) - 1)
    connector = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(marker_y + 0.27),
        Inches(right - left),
        Inches(0.05),
    )
    connector.fill.solid()
    connector.fill.fore_color.rgb = RGBColor.from_string("B6CDD9")
    connector.line.fill.background()
    text_width = min(2.15, max(1.45, spacing * 0.86))
    for index, step in enumerate(steps, start=1):
        x = left + (index - 1) * spacing
        color = BRAND_GREEN if step.get("highlight") else BRAND_BLUE
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - 0.3),
            Inches(marker_y),
            Inches(0.6),
            Inches(0.6),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor.from_string(color)
        marker.line.color.rgb = RGBColor.from_string(color)
        _set_shape_text(
            marker,
            str(index),
            size=17,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        text_x = max(0.82, min(12.48 - text_width, x - text_width / 2))
        _add_text_box(
            slide,
            str(step["title"]),
            text_x,
            3.42,
            text_width,
            0.52,
            size=15,
            color=TEXT_DARK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        description = str(step.get("description") or "").strip()
        if description:
            _add_text_box(
                slide,
                description,
                text_x,
                4.05,
                text_width,
                1.15,
                size=11.5,
                color=TEXT_MUTED,
                align=PP_ALIGN.CENTER,
            )


def _render_table(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    lead = str(spec.get("lead") or "").strip()
    table_top = 1.82 if lead else 1.48
    if lead:
        _add_text_box(slide, lead, 0.95, 1.25, 11.25, 0.42, size=15, color=TEXT_MUTED)
    columns = [str(value) for value in spec["columns"]]
    rows = [[str(value) for value in row] for row in spec["rows"]]
    table_height = 4.55 if len(rows) <= 8 else 4.72
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(columns),
        Inches(0.9),
        Inches(table_top),
        Inches(11.45),
        Inches(table_height),
    )
    table = shape.table
    widths = spec.get("column_widths") or [1.0] * len(columns)
    total = sum(float(value) for value in widths)
    for column, weight in zip(table.columns, widths):
        column.width = Inches(11.45 * float(weight) / total)
    for column_index, value in enumerate(columns):
        cell = table.cell(0, column_index)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(BRAND_BLUE)
        _set_cell_text(cell, value, size=12.5, color=WHITE, bold=True)
    highlighted = {int(value) for value in spec.get("highlight_rows") or [] if str(value).isdigit()}
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.fill.solid()
            fill = PALE_GREEN if row_index in highlighted else ("F4F7F9" if row_index % 2 else WHITE)
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            _set_cell_text(
                cell,
                value,
                size=11.5 if len(rows) <= 8 else 10.5,
                color=TEXT_DARK,
                bold=row_index in highlighted and column_index == 0,
            )


def _render_chart(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    chart_data = ChartData()
    chart_data.categories = [str(value) for value in spec["categories"]]
    for item in spec["series"]:
        chart_data.add_series(str(item["name"]), [float(value) for value in item["values"]])
    chart_type = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }[str(spec["chart_type"])]
    has_insight = bool(str(spec.get("body") or "").strip() or _strings(spec.get("bullets")))
    chart_width = 8.15 if has_insight else 11.35
    chart = slide.shapes.add_chart(
        chart_type,
        Inches(0.9),
        Inches(1.48),
        Inches(chart_width),
        Inches(4.82),
        chart_data,
    ).chart
    chart.has_title = False
    chart.has_legend = len(spec["series"]) > 1 or spec["chart_type"] in {"pie", "doughnut"}
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT_NAME
        chart.legend.font.size = Pt(10)
    chart.font.name = FONT_NAME
    chart.font.size = Pt(10)
    palette = (BRAND_BLUE, BRAND_GREEN, "5D8AA8", "6FAF8F", "8AA9BA")
    for index, series in enumerate(chart.series):
        color = RGBColor.from_string(palette[index % len(palette)])
        if spec["chart_type"] == "line":
            series.format.line.color.rgb = color
            series.format.line.width = Pt(2.25)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
            series.format.line.color.rgb = color
    if spec["chart_type"] not in {"pie", "doughnut"}:
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("D9E3E8")
        chart.value_axis.tick_labels.font.name = FONT_NAME
        chart.value_axis.tick_labels.font.size = Pt(9)
        chart.category_axis.tick_labels.font.name = FONT_NAME
        chart.category_axis.tick_labels.font.size = Pt(9)
    if has_insight:
        body = str(spec.get("body") or "").strip()
        if body:
            _add_text_box(
                slide,
                body,
                9.35,
                1.62,
                2.85,
                1.15,
                size=15,
                color=BRAND_BLUE,
                bold=True,
            )
        bullets = _strings(spec.get("bullets"))
        if bullets:
            _add_bullet_box(
                slide,
                bullets,
                9.35,
                2.95 if body else 1.65,
                2.85,
                3.0 if body else 4.25,
                size=12.5,
            )
    unit = str(spec.get("unit") or "").strip()
    if unit:
        _add_text_box(slide, f"单位：{unit}", 0.95, 1.22, 3.2, 0.24, size=9, color=TEXT_MUTED)


def _render_metrics(slide, spec: dict[str, Any]) -> None:
    _add_slide_title(slide, str(spec["title"]))
    metrics = spec["metrics"]
    count = len(metrics)
    columns = 4 if count == 4 else min(3, count)
    rows = math.ceil(count / columns)
    gap_x = 0.25
    gap_y = 0.3
    left = 0.9
    top = 2.0 if rows == 1 else 1.55
    total_width = 11.55
    total_height = 2.8 if rows == 1 else 4.75
    card_width = (total_width - gap_x * (columns - 1)) / columns
    card_height = (total_height - gap_y * (rows - 1)) / rows

    for index, metric in enumerate(metrics):
        row, column = divmod(index, columns)
        x = left + column * (card_width + gap_x)
        y = top + row * (card_height + gap_y)
        color = BRAND_BLUE if index % 2 == 0 else BRAND_GREEN
        pale = PALE_BLUE if index % 2 == 0 else PALE_GREEN
        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(card_width),
            Inches(card_height),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string(pale)
        card.line.color.rgb = RGBColor.from_string(color)
        card.line.width = Pt(1)
        _add_text_box(
            slide,
            str(metric["value"]),
            x + 0.18,
            y + 0.28,
            card_width - 0.36,
            min(0.7, card_height * 0.36),
            size=28 if rows == 1 else 24,
            color=color,
            bold=True,
        )
        _add_text_box(
            slide,
            str(metric["label"]),
            x + 0.18,
            y + min(1.05, card_height * 0.5),
            card_width - 0.36,
            0.48,
            size=14,
            color=TEXT_DARK,
            bold=True,
        )
        change = str(metric.get("change") or "").strip()
        if change:
            _add_text_box(
                slide,
                change,
                x + 0.18,
                y + card_height - 0.55,
                card_width - 0.36,
                0.32,
                size=11,
                color=TEXT_MUTED,
            )


def _render_full_image(slide, spec: dict[str, Any], spec_path: Path) -> None:
    title = str(spec.get("title") or "").strip()
    if title:
        _add_slide_title(slide, title)
    image = resolve_media_path(str(spec["image"]), spec_path)
    top = 1.45 if title else 0.85
    height = 4.85 if title else 5.45
    _add_picture(
        slide,
        image,
        0.85,
        top,
        11.65,
        height,
        fit=str(spec.get("image_fit") or "cover"),
        focus_x=_ratio(spec.get("image_focus_x"), 0.5),
        focus_y=_ratio(spec.get("image_focus_y"), 0.5),
    )
    caption = str(spec.get("caption") or "").strip()
    if caption:
        _add_text_box(
            slide,
            caption,
            0.9,
            top + height + 0.08,
            11.55,
            0.35,
            size=10,
            color=TEXT_MUTED,
            align=PP_ALIGN.CENTER,
        )


def _add_slide_title(slide, title: str) -> None:
    if len(title) <= 22:
        size = 24
    elif len(title) <= 30:
        size = 20
    else:
        size = 17
    _add_text_box(
        slide,
        title,
        3.45,
        0.38,
        7.25,
        0.72,
        size=size,
        color=TEXT_DARK,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def _add_column(
    slide,
    payload: dict[str, Any],
    x: float,
    y: float,
    width: float,
    height: float,
    color: str,
    pale: str,
) -> None:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string(pale)
    panel.line.color.rgb = RGBColor.from_string(color)
    panel.line.width = Pt(1)
    heading = str(payload.get("title") or "").strip()
    if heading:
        _add_text_box(
            slide, heading, x + 0.28, y + 0.25, width - 0.56, 0.55, size=18, color=color, bold=True
        )
    bullets = _strings(payload.get("bullets"))
    body = str(payload.get("body") or "").strip()
    cursor = y + 0.95
    if body:
        _add_text_box(slide, body, x + 0.28, cursor, width - 0.56, 1.0, size=15, color=TEXT_DARK)
        cursor += 1.12
    if bullets:
        _add_bullet_box(
            slide,
            bullets,
            x + 0.28,
            cursor,
            width - 0.56,
            max(0.8, y + height - cursor - 0.25),
            size=15,
        )


def _add_comparison_panel(slide, payload: dict[str, Any], x: float, color: str, pale: str) -> None:
    width = 5.55
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.55), Inches(width), Inches(0.72)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor.from_string(color)
    header.line.color.rgb = RGBColor.from_string(color)
    _set_shape_text(
        header,
        str(payload.get("title") or ""),
        size=18,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    body = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.27), Inches(width), Inches(4.0)
    )
    body.fill.solid()
    body.fill.fore_color.rgb = RGBColor.from_string(pale)
    body.line.color.rgb = RGBColor.from_string(color)
    body.line.width = Pt(1)
    bullets = _strings(payload.get("bullets"))
    body_text = str(payload.get("body") or "").strip()
    cursor = 2.55
    if body_text:
        _add_text_box(slide, body_text, x + 0.3, cursor, width - 0.6, 1.05, size=15, color=TEXT_DARK)
        cursor += 1.15
    if bullets:
        _add_bullet_box(slide, bullets, x + 0.3, cursor, width - 0.6, 6.0 - cursor, size=15)


def _add_text_box(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    _set_shape_text(
        shape,
        text,
        size=size,
        color=color,
        bold=bold,
        align=align,
        valign=valign,
    )
    return shape


def _set_shape_text(
    shape,
    text: str,
    *,
    size: float,
    color: str,
    bold: bool,
    align,
    valign,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = valign
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    _style_paragraph(paragraph, size=size, color=color, bold=bold)


def _add_bullet_box(
    slide,
    bullets: Iterable[str],
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float,
):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 0
        paragraph.space_after = Pt(9)
        paragraph.line_spacing = 1.15
        _set_real_bullet(paragraph)
        _style_paragraph(paragraph, size=size, color=TEXT_DARK, bold=False)
    return shape


def _set_real_bullet(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar"):
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "•")
    properties.insert(0, bullet)


def _style_paragraph(paragraph, *, size: float, color: str, bold: bool) -> None:
    for run in paragraph.runs:
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_text(cell, text: str, *, size: float, color: str, bold: bool) -> None:
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.LEFT
    _style_paragraph(paragraph, size=size, color=color, bold=bold)


def _add_picture(
    slide,
    image_path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    fit: str = "cover",
    focus_x: float = 0.5,
    focus_y: float = 0.5,
):
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    if image_width <= 0 or image_height <= 0:
        raise DeckSpecError(f"image has invalid dimensions: {image_path}")
    image_ratio = image_width / image_height
    frame_ratio = width / height
    if fit == "contain":
        if image_ratio > frame_ratio:
            rendered_width = width
            rendered_height = width / image_ratio
        else:
            rendered_height = height
            rendered_width = height * image_ratio
        return slide.shapes.add_picture(
            str(image_path),
            Inches(x + (width - rendered_width) / 2),
            Inches(y + (height - rendered_height) / 2),
            width=Inches(rendered_width),
            height=Inches(rendered_height),
        )
    if fit != "cover":
        raise DeckSpecError(f"unsupported image fit: {fit}")
    picture = slide.shapes.add_picture(
        str(image_path),
        Inches(x),
        Inches(y),
        width=Inches(width),
        height=Inches(height),
    )
    if image_ratio > frame_ratio:
        total_crop = 1 - frame_ratio / image_ratio
        picture.crop_left = total_crop * focus_x
        picture.crop_right = total_crop * (1 - focus_x)
    elif image_ratio < frame_ratio:
        total_crop = 1 - image_ratio / frame_ratio
        picture.crop_top = total_crop * focus_y
        picture.crop_bottom = total_crop * (1 - focus_y)
    return picture


def _ratio(value: Any, default: float) -> float:
    if not isinstance(value, (int, float)):
        return default
    return min(1.0, max(0.0, float(value)))


def _add_footer(
    slide,
    deck_spec: dict[str, Any],
    slide_spec: dict[str, Any],
    page_number: int,
) -> None:
    footer = str(slide_spec.get("footer") or deck_spec.get("footer") or "").strip()
    source = str(slide_spec.get("source") or "").strip()
    footer_parts = [part for part in (footer, f"来源：{source}" if source else "") if part]
    if footer_parts:
        _add_text_box(
            slide,
            "  |  ".join(footer_parts),
            0.85,
            6.45,
            10.45,
            0.3,
            size=8,
            color=TEXT_MUTED,
        )
    if deck_spec.get("show_page_numbers", True):
        _add_text_box(
            slide,
            str(page_number),
            11.65,
            6.45,
            0.5,
            0.28,
            size=9,
            color=TEXT_MUTED,
            align=PP_ALIGN.RIGHT,
        )


def _strings(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]
