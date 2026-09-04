#!/usr/bin/env python3
"""Validate a generated corporate PPTX structurally and semantically."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from pptx import Presentation

from ppt_common import DEFAULT_TEMPLATE, DeckSpecError, load_deck_spec

PLACEHOLDER_MARKERS = (
    "单击此处",
    "按一下以",
    "click to add",
    "lorem ipsum",
    "[insert",
)
FORBIDDEN_PART_PREFIXES = (
    "ppt/activeX/",
    "ppt/vbaProject",
)
ALLOWED_CHART_WORKBOOK = re.compile(
    r"^ppt/embeddings/Microsoft_Excel_(?:Sheet|Worksheet)\d*\.xlsx$"
)


def validate(
    deck_path: Path,
    *,
    spec_path: Path | None = None,
    template_path: Path = DEFAULT_TEMPLATE,
) -> dict:
    deck = deck_path.expanduser().resolve()
    template = template_path.expanduser().resolve()
    errors: list[str] = []
    warnings: list[str] = []

    if not deck.is_file():
        raise DeckSpecError(f"PowerPoint file not found: {deck}")
    if deck.suffix.lower() != ".pptx":
        raise DeckSpecError("validation input must end in .pptx")
    if not template.is_file():
        raise DeckSpecError(f"PowerPoint template not found: {template}")

    try:
        with zipfile.ZipFile(deck) as archive:
            bad_member = archive.testzip()
            if bad_member:
                errors.append(f"corrupt ZIP member: {bad_member}")
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "ppt/presentation.xml"}
            for missing in sorted(required - names):
                errors.append(f"missing required package part: {missing}")
            for name in names:
                if name.startswith(FORBIDDEN_PART_PREFIXES):
                    errors.append(f"forbidden active content or embedded object: {name}")
                if name.startswith("ppt/embeddings/"):
                    _check_chart_workbook(archive, name, errors)
                if name.endswith(".rels"):
                    _check_relationships(archive, name, errors)
    except zipfile.BadZipFile as exc:
        errors.append(f"invalid PPTX ZIP package: {exc}")

    presentation = None
    template_presentation = None
    try:
        presentation = Presentation(deck)
        template_presentation = Presentation(template)
    except Exception as exc:  # python-pptx provides several package-specific exceptions
        errors.append(f"python-pptx could not open the deck: {exc}")

    if presentation is not None and template_presentation is not None:
        if (presentation.slide_width, presentation.slide_height) != (
            template_presentation.slide_width,
            template_presentation.slide_height,
        ):
            errors.append("slide dimensions do not match the corporate template")
        if len(presentation.slide_masters) < len(template_presentation.slide_masters):
            errors.append("one or more corporate slide masters were removed")
        _check_slide_text(presentation, errors, warnings)

        if spec_path is not None:
            spec = load_deck_spec(spec_path.expanduser().resolve())
            expected = len(spec["slides"])
            actual = len(presentation.slides)
            if actual != expected:
                errors.append(f"slide count mismatch: expected {expected}, found {actual}")

    return {
        "ok": not errors,
        "deck": str(deck),
        "slides": len(presentation.slides) if presentation is not None else None,
        "errors": errors,
        "warnings": warnings,
    }


def _check_relationships(archive: zipfile.ZipFile, name: str, errors: list[str]) -> None:
    try:
        body = archive.read(name)
        if len(body) > 4 * 1024 * 1024:
            errors.append(f"relationship part is unexpectedly large: {name}")
            return
        root = ElementTree.fromstring(body)
    except (KeyError, ElementTree.ParseError) as exc:
        errors.append(f"invalid relationship part {name}: {exc}")
        return
    for relationship in root:
        if relationship.attrib.get("TargetMode", "").lower() == "external":
            target = relationship.attrib.get("Target", "")
            errors.append(f"external relationship is not allowed: {name} -> {target}")


def _check_chart_workbook(archive: zipfile.ZipFile, name: str, errors: list[str]) -> None:
    if ALLOWED_CHART_WORKBOOK.fullmatch(name) is None:
        errors.append(f"unsupported embedded object: {name}")
        return
    try:
        info = archive.getinfo(name)
    except KeyError:
        errors.append(f"embedded chart workbook is missing: {name}")
        return
    if info.file_size > 5 * 1024 * 1024:
        errors.append(f"embedded chart workbook is unexpectedly large: {name}")


def _check_slide_text(presentation, errors: list[str], warnings: list[str]) -> None:
    for index, slide in enumerate(presentation.slides, start=1):
        visible_text = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text.strip()
            if not text:
                continue
            visible_text.append(text)
            lowered = text.lower()
            if any(marker in lowered for marker in PLACEHOLDER_MARKERS):
                errors.append(f"slide {index} contains leftover placeholder text: {text[:80]}")
            if len(text) > 800:
                warnings.append(f"slide {index} contains a text box longer than 800 characters")
        if index not in {1, len(presentation.slides)} and not visible_text:
            warnings.append(f"slide {index} contains no editable text")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Validate a generated corporate PPTX.")
    parser.add_argument("deck", type=Path, help="generated .pptx file")
    parser.add_argument("--spec", type=Path, help="matching CorporateDeck YAML or JSON")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument(
        "--strict",
        action="store_true",
        help="treat warnings as validation failures",
    )
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        result = validate(args.deck, spec_path=args.spec, template_path=args.template)
    except (DeckSpecError, OSError, ValueError) as exc:
        result = {"ok": False, "errors": [str(exc)], "warnings": []}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["ok"] or (args.strict and result.get("warnings")):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
