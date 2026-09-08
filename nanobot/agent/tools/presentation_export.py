"""Export one gateway-bound presentation using its selected renderer."""

from __future__ import annotations

import asyncio
import base64
import json
import mimetypes
import os
import re
import signal
import sys
import tempfile
from html import escape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit

from nanobot.agent.tools.base import tool_parameters
from nanobot.agent.tools.context import current_request_session_key
from nanobot.agent.tools.filesystem import _FsTool
from nanobot.agent.tools.presentation import _artifact
from nanobot.agent.tools.schema import StringSchema, tool_parameters_schema
from nanobot.observability.operations import Operation, operation
from nanobot.presentations import (
    PresentationError,
    PresentationService,
    runtime_executable,
    source_digest,
    template_by_id,
)


class _Slides(HTMLParser):
    def __init__(self):
        super().__init__()
        self.count = 0

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if tag == "section" and "slide" in (values.get("class") or "").split():
            self.count += 1


class _PortableHTML(HTMLParser):
    """Embed project-relative assets so srcDoc previews and shared HTML keep their media."""

    def __init__(self, folder: Path):
        super().__init__(convert_charrefs=False)
        self.folder = folder
        self.parts: list[str] = []
        self.in_style = False
        self.in_script = False
        self.bytes = 0

    def asset(self, url: str) -> str:
        parsed = urlsplit(url)
        if parsed.scheme in {"https", "http", "data", "blob"} or url.startswith(("#", "//")):
            return url
        if parsed.scheme or not parsed.path:
            raise PresentationError(f"Unsupported presentation asset: {url}")
        path = (self.folder / unquote(parsed.path)).resolve()
        if not path.is_relative_to(self.folder) or ".source" in path.relative_to(self.folder).parts:
            raise PresentationError("Presentation assets must be inside the document project")
        size = path.stat().st_size
        self.bytes += size
        if size > 25 * 1024 * 1024 or self.bytes > 100 * 1024 * 1024:
            raise PresentationError("Presentation media exceeds the export size limit")
        mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode()

    def css(self, text: str) -> str:
        return re.sub(r"url\(\s*(['\"]?)([^)\s]+?)\1\s*\)",
                      lambda match: f'url("{self.asset(match[2])}")', text)

    def handle_starttag(self, tag, attrs):
        if tag == "base":
            raise PresentationError("Presentation HTML must not override its base URL")
        values = dict(attrs)
        if "srcset" in values:
            raise PresentationError("Use one local src per presentation image instead of srcset")
        for key, value in list(values.items()):
            if value and (key in {"src", "poster"} or (tag == "link" and key == "href")):
                values[key] = self.asset(value)
            elif key == "style" and value:
                values[key] = self.css(value)
        self.parts.append("<" + tag + "".join(
            f' {key}="{escape(value, quote=True)}"' if value is not None else f" {key}"
            for key, value in values.items()
        ) + ">")
        self.in_style = tag == "style" or self.in_style
        self.in_script = tag == "script" or self.in_script

    def handle_endtag(self, tag):
        self.parts.append(f"</{tag}>")
        if tag == "style":
            self.in_style = False
        if tag == "script":
            self.in_script = False

    def handle_data(self, data):
        if self.in_script and "'./assets/motion.min.js'" in data:
            data = data.replace("'./assets/motion.min.js'", json.dumps(self.asset("assets/motion.min.js")))
        self.parts.append(self.css(data) if self.in_style else data)

    def handle_comment(self, data):
        self.parts.append(f"<!--{data}-->")

    def handle_decl(self, decl):
        self.parts.append(f"<!{decl}>")

    def handle_entityref(self, name):
        self.parts.append(f"&{name};")

    def handle_charref(self, name):
        self.parts.append(f"&#{name};")


async def _script(arguments: list[str], cwd: Path, *, timeout: int = 300) -> str:
    stage = Path(arguments[1]).name if len(arguments) > 1 and arguments[1] != "-m" else "kimi_render"
    with operation("presentation.script", stage=stage, timeout_ms=timeout * 1000):
        return await _run_script(arguments, cwd, timeout=timeout)


async def _run_script(arguments: list[str], cwd: Path, *, timeout: int = 300) -> str:
    with tempfile.TemporaryFile() as log:
        process = await asyncio.create_subprocess_exec(
            *arguments, cwd=str(cwd), stdout=log, stderr=asyncio.subprocess.STDOUT,
            start_new_session=os.name != "nt",
            env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1",
                 "PYTHONDONTWRITEBYTECODE": "1"},
        )
        try:
            await asyncio.wait_for(process.wait(), timeout)
        except BaseException:
            if process.returncode is None:
                if os.name == "nt":
                    process.kill()
                else:
                    os.killpg(process.pid, signal.SIGKILL)
            await process.wait()
            raise
        log.seek(max(0, log.tell() - 16000))
        text = log.read().decode("utf-8", errors="replace")
    if process.returncode:
        raise PresentationError(text[-4000:] or "Presentation export failed")
    return text


def _preview_input(presentation, path: Path) -> Path:
    """Quick Look rejects some native notes masters; notes do not affect slide previews."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    for slide in presentation.slides:
        for relationship in list(slide.part.rels.values()):
            if relationship.reltype == RT.NOTES_SLIDE:
                slide.part.drop_rel(relationship.rId)
    for relationship in list(presentation.part.rels.values()):
        if relationship.reltype == RT.NOTES_MASTER:
            presentation.part.drop_rel(relationship.rId)
    for node in presentation._element.xpath("./p:notesMasterIdLst"):
        presentation._element.remove(node)
    presentation.save(path)
    return path


@tool_parameters(tool_parameters_schema(
    document_id=StringSchema("Document ID assigned by the presentation picker.", min_length=8),
    required=["document_id"],
))
class ExportPresentationTool(_FsTool):
    _scopes = {"core"}

    @property
    def name(self) -> str:
        return "export_presentation"

    @property
    def description(self) -> str:
        return (
            "Export the explicitly selected presentation document after writing its sources. "
            "Uses the document's pinned Taiping PPTX, Guizang HTML, or Kimi PPTD/PPTX renderer. "
            "For single-page edits change only the requested source page, then export this same ID. "
            "Never install dependencies or use another renderer to bypass an export error."
        )

    async def execute(self, document_id: str, **kwargs: Any) -> str | dict[str, Any]:
        with operation("presentation.export", document_id=document_id, stage="preflight") as observed:
            return await self._execute(document_id, observed)

    async def _execute(self, document_id: str, observed: Operation) -> str | dict[str, Any]:
        try:
            service = PresentationService(self._workspace or Path.cwd())
            session_key = current_request_session_key()
            if not session_key:
                observed.fail("PRESENTATION_CONTEXT_REQUIRED")
                return "Error: presentation export requires a conversation context"
            document = service.document(document_id, session_key)
            template = template_by_id(document["template_id"])
            observed.details["template_id"] = template["id"]
            folder = self._resolve_write(document["project_path"])
            source = self._resolve_read(str(folder / ".source"))
            if source_digest(source) != document["source_digest"]:
                observed.fail("PRESENTATION_SNAPSHOT_CHANGED")
                return "Error: the bound template snapshot has changed; create a new document"
            if any(path.is_symlink() for path in folder.rglob("*")):
                observed.fail("PRESENTATION_SYMLINK_REJECTED")
                return "Error: presentation projects must not contain symlinks"
            missing = await asyncio.to_thread(service.availability, template, source)
            if missing:
                observed.fail("PRESENTATION_DEPENDENCY_MISSING")
                return "Error: presentation unavailable: " + ", ".join(missing)
            output = self._resolve_write(str(folder / f"presentation.{template['format']}"))
            warnings = []
            # Publish only validated output; failed revisions preserve the last successful export.
            with tempfile.TemporaryDirectory(prefix=".export-", dir=folder) as staging:
                staged = Path(staging) / output.name
                observed.details["stage"] = "render"
                with operation("presentation.render", template_id=template["id"]):
                    artifacts, page_count, warnings = await self._export(template, folder, source, staged)
                observed.details["stage"] = "publish"
                with operation("storage.artifact_publish", document_id=document_id):
                    os.replace(staged, output)
                artifacts.insert(0, _artifact(output, "text/html" if template["format"] == "html" else
                                             "application/vnd.openxmlformats-officedocument.presentationml.presentation"))
            observed.details["stage"] = "register"
            with operation("storage.artifact_register", document_id=document_id, page_count=page_count):
                service.complete(document_id, artifacts=artifacts, page_count=page_count)
            observed.details.update({"page_count": page_count, "warning_count": len(warnings), "artifact_count": len(artifacts)})
            return {"text": json.dumps({"document_id": document_id, "template_id": template["id"],
                                        "page_count": page_count, "project_path": str(folder),
                                        "warnings": warnings}, ensure_ascii=False),
                    "files": artifacts}
        except (PresentationError, OSError, ValueError, asyncio.TimeoutError) as exc:
            observed.fail("PRESENTATION_EXPORT_FAILED")
            observed.details["error_type"] = type(exc).__name__
            return f"Error: presentation export failed: {exc}"

    async def _export(self, template, folder, source, output):
        artifacts = []
        warnings = []
        if template["family"] == "taiping":
            spec = self._resolve_read(str(folder / "deck.yaml"))
            pinned = str(source / "assets/ppt-template.pptx")
            await _script([sys.executable, str(source / "scripts/generate_deck.py"), str(spec),
                           "--output", str(output), "--template", pinned], folder)
            validation = await _script([sys.executable, str(source / "scripts/validate_deck.py"),
                                        str(output), "--spec", str(spec), "--template", pinned], folder)
            warnings.extend(json.loads(validation).get("warnings", []))
            from pptx import Presentation
            page_count = len(Presentation(output).slides)
            preview = self._resolve_write(str(folder / "preview.pdf"))
            try:
                await _script([sys.executable, str(source / "scripts/render_preview.py"),
                               str(output), "--output", str(preview), "--force"], folder)
                artifacts.append(_artifact(preview, "application/pdf"))
            except (PresentationError, asyncio.TimeoutError) as exc:
                warnings.append(f"PDF preview unavailable: {exc}")
        elif template["family"] == "guizang":
            html_path = self._resolve_read(str(folder / "index.html"))
            html = html_path.read_text(encoding="utf-8")
            parser = _Slides()
            parser.feed(html)
            if not parser.count or "SLIDES_HERE" in html or "[必填]" in html:
                raise PresentationError("Fill the presentation slides and title before exporting")
            if template["id"] == "guizang-swiss":
                await _script([runtime_executable("node"), str(source / "scripts/validate-swiss-deck.mjs"), str(html_path)], folder)
            portable = _PortableHTML(folder)
            portable.feed(html)
            portable.close()
            output.write_text("".join(portable.parts), encoding="utf-8")
            page_count = parser.count
        else:
            manifest = self._resolve_read(str(folder / "deck.pptd"))
            await _script([
                sys.executable, "-m", "nanobot.presentation_kimi",
                str(manifest), str(output),
            ], folder)
            from pptx import Presentation
            presentation = Presentation(output)
            page_count = len(presentation.slides)
            artifacts.append(_artifact(manifest, "application/yaml"))
            from nanobot.agent.skills import BUILTIN_SKILLS_DIR
            preview_script = source / "scripts/render_preview.py"
            if not preview_script.is_file():
                preview_script = BUILTIN_SKILLS_DIR / "corporate-ppt/scripts/render_preview.py"
            preview = self._resolve_write(str(folder / "preview.pdf"))
            staged_preview = output.parent / "preview.pdf"
            try:
                preview_deck = _preview_input(presentation, output.parent / "preview-input.pptx")
                await _script([sys.executable, str(preview_script), str(preview_deck),
                               "--output", str(staged_preview), "--force"], folder)
                os.replace(staged_preview, preview)
                artifacts.append(_artifact(preview, "application/pdf"))
            except (PresentationError, asyncio.TimeoutError) as exc:
                warnings.append(f"PDF preview unavailable: {exc}")
        if page_count < 1:
            raise PresentationError("Presentation contains no pages")
        return artifacts, page_count, warnings
