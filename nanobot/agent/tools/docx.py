"""Generate Word documents using the user's bundled document template."""

from __future__ import annotations

import asyncio
import re
from pathlib import Path
from typing import Any

from nanobot.agent.tools.base import tool_parameters
from nanobot.agent.tools.filesystem import FileToolsConfig, _FsTool
from nanobot.agent.tools.pdf import _markdown_blocks
from nanobot.agent.tools.schema import StringSchema, tool_parameters_schema

_DEFAULT_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2] / "skills" / "office-documents" / "assets" / "template.docx"
)


@tool_parameters(
    tool_parameters_schema(
        content=StringSchema("Document text, with optional Markdown formatting. Prefer this to creating a source file. Supply content or source_path, not both.", min_length=1, nullable=True),
        source_path=StringSchema("Optional existing Markdown/plain-text source file, instead of content.", min_length=1, nullable=True),
        output_path=StringSchema("Use the output_path returned by prepare_output. Required with content; defaults to a versioned source filename otherwise. Always use the actual returned file path.", nullable=True),
        title=StringSchema("Optional document title. Defaults to the first H1 or source filename.", nullable=True),
        template_path=StringSchema("Optional user-requested DOCX template. Uses the bundled user template by default. Replaces sample body text, preserving styles, page setup, headers and footers.", nullable=True),
        template=StringSchema("Legacy compatibility argument; named styles such as research_report no longer replace the user template.", nullable=True),
        classification=StringSchema("Optional confidentiality label explicitly provided by the user; never infer one from the template example.", nullable=True),
        recipient=StringSchema("Optional recipient provided by the user.", nullable=True),
        signatory=StringSchema("Optional signing organization or person provided by the user.", nullable=True),
        document_date=StringSchema("Optional document date provided by the user.", nullable=True),
        required=[],
    )
)
class CreateDocxTool(_FsTool):
    """Create a Word artifact without requiring a Markdown companion."""

    _scopes = {"core", "subagent"}
    config_key = "file"

    @classmethod
    def config_cls(cls):
        return FileToolsConfig

    @property
    def name(self) -> str:
        return "create_docx"

    @property
    def description(self) -> str:
        return (
            "Create a DOCX using the bundled user-provided Word template. Prefer inline content "
            "and output_path; no Markdown or HTML source/companion files are created. "
            "An existing source_path is also supported. Returns the DOCX as a structured artifact. "
            "Do not substitute a self-designed template or install conversion dependencies."
        )

    async def execute(
        self,
        source_path: str | None = None,
        output_path: str | None = None,
        title: str | None = None,
        template: str | None = None,
        content: str | None = None,
        template_path: str | None = None,
        classification: str | None = None,
        recipient: str | None = None,
        signatory: str | None = None,
        document_date: str | None = None,
        **kwargs: Any,
    ) -> str | dict[str, Any]:
        if (content is None) == (not source_path):
            return self._error("render_failed", "provide exactly one of content or source_path", source_path or "")
        if content is not None and not output_path:
            return self._error("render_failed", "output_path is required with content", "")

        try:
            output = self._resolve_write(output_path or str(Path(source_path).with_suffix(".docx")))
            source = self._resolve_read(source_path) if source_path else output
            word_template = self._resolve_read(template_path) if template_path else _DEFAULT_TEMPLATE_PATH
        except Exception as exc:
            return self._error("permission_denied", str(exc), source_path or "")

        if output.suffix.lower() != ".docx":
            return self._error("render_failed", "output_path must end in .docx", source_path or "")
        if output == word_template.resolve():
            return self._error("render_failed", "output_path must not overwrite the Word template", source_path or "")
        if source_path:
            if source.suffix.lower() not in {".md", ".markdown", ".txt"}:
                return self._error("render_failed", "source_path must be a Markdown or text file", str(source))
            try:
                content = source.read_text(encoding="utf-8")
            except Exception as exc:
                return self._error("render_failed", f"failed to read source: {exc}", str(source))
        if not content.strip():
            return self._error("render_failed", "document content is empty", source_path or "")
        if not word_template.is_file():
            return self._error("render_failed", "Word template not found", str(word_template))

        render_title = (title or _title_from_markdown(content, source)).strip()
        try:
            output = self._versioned_output_path(output)
            staged = self._staged_output_path(output)
            await asyncio.to_thread(
                _render_docx, content, source, staged, render_title, word_template,
                classification=classification, recipient=recipient,
                signatory=signatory, document_date=document_date,
            )
        except ModuleNotFoundError as exc:
            return self._error("dependency_missing", f"missing dependency: {exc.name}", str(source))
        except Exception as exc:
            return self._error("render_failed", str(exc), str(source))

        valid, validation_error = _validate_docx(staged)
        if not valid:
            return self._error("validation_failed", validation_error, str(source))
        try:
            self._publish_output(staged, output)
        except (OSError, ValueError) as exc:
            return self._error("publish_failed", str(exc), str(source))
        size = output.stat().st_size
        return {
            "text": (
                "DOCX created successfully\n"
                + (f"source_path: {source}\n" if source_path else "")
                + f"template_path: {word_template}\n"
                + f"docx_path: {output}\n"
                f"file_size: {size}"
            ),
            "files": [{
                "path": str(output),
                "name": output.name,
                "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "size": size,
            }],
        }

    def _error(self, code: str, message: str, source_path: str) -> str:
        return f"Error: {code}: {message}\nsource_path: {source_path}"


def _title_from_markdown(content: str, source: Path) -> str:
    match = re.search(r"^\s*#\s+(.+?)\s*$", content, flags=re.MULTILINE)
    if match:
        return _plain_markdown(match.group(1)) or source.stem
    return source.stem.replace("_", " ").replace("-", " ").strip() or "Document"


def _plain_markdown(value: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", value)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"__([^_]+)__", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"_([^_]+)_", r"\1", text)
    return text.strip()


_INLINE_TOKEN_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|`.+?`|\*[^*]+?\*|_[^_]+?_)")


def _set_run_text(run: Any, value: str) -> None:
    run.add_text(_plain_markdown(value))


def _add_inline(paragraph: Any, value: str) -> None:
    """Render a conservative Markdown inline subset without external conversion tools."""
    for token in _INLINE_TOKEN_RE.split(value):
        if not token:
            continue
        run = paragraph.add_run()
        if token.startswith("**") and token.endswith("**"):
            run.bold = True
            _set_run_text(run, token[2:-2])
        elif token.startswith("__") and token.endswith("__"):
            run.bold = True
            _set_run_text(run, token[2:-2])
        elif token.startswith("`") and token.endswith("`"):
            run.font.name = "Menlo"
            _set_run_text(run, token[1:-1])
        elif token.startswith("*") and token.endswith("*"):
            run.italic = True
            _set_run_text(run, token[1:-1])
        elif token.startswith("_") and token.endswith("_"):
            run.italic = True
            _set_run_text(run, token[1:-1])
        else:
            _set_run_text(run, token)


def _add_paragraph(document: Any, value: str = "", *, style: str = "Body Text") -> Any:
    paragraph = document.add_paragraph(style=style if style in document.styles else "Normal")
    if value:
        _add_inline(paragraph, value)
    return paragraph


def _add_table(document: Any, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    if "Table Grid" in document.styles:
        table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(column_count):
            cell = cells[column_index]
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            _add_inline(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def _add_image(document: Any, source: Path, image_ref: str, caption: str) -> bool:
    candidate = Path(image_ref.strip().strip('"\''))
    try:
        if candidate.is_absolute():
            return False
        source_directory = source.parent.resolve()
        resolved = (source_directory / candidate).resolve()
        resolved.relative_to(source_directory)
        if not resolved.is_file():
            return False
        document.add_picture(str(resolved))
        if caption:
            paragraph = _add_paragraph(document, style="Caption")
            paragraph.alignment = 1
            paragraph.add_run(caption)
        return True
    except Exception:
        return False


def _render_docx(
    content: str, source: Path, output: Path, title: str, template_path: Path,
    *, classification: str | None = None, recipient: str | None = None,
    signatory: str | None = None, document_date: str | None = None,
) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    # Fail clearly if the user's template is unavailable; never substitute Document().
    document = Document(str(template_path))
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    document.core_properties.title = title

    if classification:
        _add_paragraph(document, classification, style="Classification")
    _add_paragraph(document, title, style="Title")
    if recipient:
        _add_paragraph(document, recipient, style="Recipient")

    for kind, value in _markdown_blocks(content):
        if re.fullmatch(r"h[1-6]", kind):
            level = int(kind[1:])
            heading_text = _plain_markdown(value)
            if level == 1 and heading_text == title:
                continue
            _add_paragraph(document, heading_text, style=f"Heading {min(3, max(1, level - 1))}")
        elif kind == "image":
            caption, image_ref = value
            _add_image(document, source, image_ref, caption)
        elif kind == "paragraph":
            _add_paragraph(document, value)
        elif kind == "bullet":
            _add_paragraph(document, "• " + value)
        elif kind == "numbered":
            number, item = value
            _add_paragraph(document, f"{number}. {item}")
        elif kind == "blockquote":
            paragraph = _add_paragraph(document)
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.right_indent = Cm(0.4)
            _add_inline(paragraph, value)
            for run in paragraph.runs:
                run.italic = True
        elif kind == "hr":
            _add_paragraph(document, "—" * 30).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "code":
            paragraph = _add_paragraph(document)
            paragraph.paragraph_format.left_indent = Cm(0.5)
            run = paragraph.add_run(value)
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
        elif kind == "mermaid":
            paragraph = _add_paragraph(document)
            paragraph.add_run("图示源：\n").bold = True
            run = paragraph.add_run(value)
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
        elif kind == "table":
            _add_table(document, value)

    if signatory:
        _add_paragraph(document, signatory, style="Signature")
    if document_date:
        _add_paragraph(document, document_date, style="Document Date")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def _validate_docx(path: Path) -> tuple[bool, str]:
    if not path.exists():
        return False, "DOCX file was not created"
    size = path.stat().st_size
    if size < 4_000:
        return False, f"DOCX file is too small ({size} bytes)"
    try:
        from docx import Document

        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    except Exception as exc:
        return False, f"DOCX cannot be opened: {exc}"
    if not text:
        return False, "DOCX contains no readable text"
    return True, ""
