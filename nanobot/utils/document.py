"""Document text extraction utilities for nanobot."""

import json
import mimetypes
import platform
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from loguru import logger

from nanobot.utils.helpers import detect_image_mime

# Supported file extensions for text extraction
SUPPORTED_EXTENSIONS: set[str] = {
    # Document formats
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".pptx",
    # Text formats
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    # Image formats (for future OCR support)
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
}

_MAX_TEXT_LENGTH = 200_000
_WORD_CONVERSION_TIMEOUT_SECONDS = 45


def extract_text(path: Path) -> str | None:
    """Extract text from a file.

    Args:
        path: Path to the file.

    Returns:
        Extracted text as string, None for unsupported types,
        or error string for failures.
    """
    if not isinstance(path, Path):
        path = Path(path)

    if not path.exists():
        return f"[error: file not found: {path}]"

    ext = path.suffix.lower()

    # Document formats -- each branch lazily imports its parser so that
    # startup does not pay the ~25 MB cost of loading openpyxl /
    # python-docx / python-pptx / pypdf up front (see issue #3422).
    if ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".doc":
        return _extract_doc(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext == ".xls":
        return _extract_xls(path)
    elif ext == ".xlsx":
        return _extract_xlsx(path)
    elif ext == ".pptx":
        return _extract_pptx(path)
    elif _is_text_extension(ext):
        return _extract_text_file(path)
    elif ext in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        # Image files - for future OCR support
        return f"[image: {path.name}]"
    else:
        # Unsupported extension
        return None


def _extract_pdf(path: Path) -> str:
    """Extract text from PDF using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[error: pypdf not installed]"
    try:
        reader = PdfReader(path)
        pages: list[str] = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            pages.append(f"--- Page {i} ---\n{text}")
        return _truncate("\n\n".join(pages), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract PDF {}", path)
        return f"[error: failed to extract PDF: {e!s}]"


def _extract_docx(path: Path) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document as DocxDocument
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return "[error: python-docx not installed]"
    try:
        doc = DocxDocument(path)
        blocks: list[str] = []
        paragraph_tag = qn("w:p")
        table_tag = qn("w:tbl")
        for child in doc.element.body.iterchildren():
            if child.tag == paragraph_tag:
                text = Paragraph(child, doc).text
                if text.strip():
                    blocks.append(text)
            elif child.tag == table_tag:
                rows: list[str] = []
                for row in Table(child, doc).rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    row_text = "\t".join(cells).rstrip()
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    blocks.append("\n".join(rows))
        return _truncate("\n\n".join(blocks), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract DOCX {}", path)
        return f"[error: failed to extract DOCX: {e!s}]"


def _extract_doc(path: Path) -> str:
    """Convert a legacy DOC with desktop Word, then use the DOCX extractor."""
    if platform.system() != "Windows":
        return (
            "[error: legacy .doc reading requires Microsoft Word desktop on Windows; "
            "convert the file to .docx before reading it on this platform]"
        )

    with tempfile.TemporaryDirectory(
        prefix="nanobot-word-", ignore_cleanup_errors=True
    ) as temporary:
        converted = Path(temporary) / "converted.docx"
        error = _convert_doc_with_word(path, converted)
        if error is not None:
            return error
        return _extract_docx(converted)


def _convert_doc_with_word(source: Path, destination: Path) -> str | None:
    command = [
        sys.executable,
        "-m",
        "nanobot.utils._word_com_worker",
        str(source.resolve()),
        str(destination.resolve()),
    ]
    run_options: dict[str, Any] = {
        "capture_output": True,
        "check": False,
        "encoding": "utf-8",
        "errors": "replace",
        "text": True,
        "timeout": _WORD_CONVERSION_TIMEOUT_SECONDS,
    }
    if platform.system() == "Windows":
        run_options["creationflags"] = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    try:
        completed = subprocess.run(command, **run_options)
    except subprocess.TimeoutExpired:
        logger.warning(
            "Microsoft Word conversion timed out after {} seconds for {}",
            _WORD_CONVERSION_TIMEOUT_SECONDS,
            source,
        )
        return (
            "[error: Microsoft Word timed out while converting the legacy .doc file "
            f"after {_WORD_CONVERSION_TIMEOUT_SECONDS} seconds]"
        )
    except OSError as exc:
        logger.warning("Could not start Microsoft Word conversion for {}: {}", source, exc)
        return f"[error: could not start Microsoft Word conversion: {exc!s}]"

    payload = _word_worker_payload(completed.stdout)
    if completed.returncode != 0:
        code = str(payload.get("code") or "WORD_CONVERSION_FAILED")
        message = str(
            payload.get("message")
            or completed.stderr.strip()
            or "Microsoft Word could not convert the legacy .doc file"
        )
        detail = str(payload.get("detail") or "")
        logger.warning(
            "Microsoft Word conversion failed for {}: {} {} {}",
            source,
            code,
            message,
            detail,
        )
        return f"[error: {code}: {message}]"

    if not destination.is_file() or destination.stat().st_size == 0:
        logger.warning("Microsoft Word reported success without producing {}", destination)
        return "[error: WORD_OUTPUT_MISSING: Microsoft Word produced no converted .docx file]"
    return None


def _word_worker_payload(stdout: str) -> dict[str, Any]:
    for line in reversed(stdout.splitlines()):
        try:
            value = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            return value
    return {}


def _extract_xls(path: Path) -> str:
    """Extract cell values from a legacy XLS workbook using xlrd."""
    try:
        import xlrd
    except ImportError:
        return "[error: xlrd not installed]"

    workbook = None
    try:
        workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
        sheets: list[str] = []
        for sheet in workbook.sheets():
            rows: list[str] = []
            for row_index in range(sheet.nrows):
                cells = [
                    _format_xls_cell(cell, workbook, xlrd)
                    for cell in sheet.row(row_index)
                ]
                row_text = "\t".join(cells).rstrip()
                if row_text.strip():
                    rows.append(row_text)
            if rows:
                sheets.append(f"--- Sheet: {sheet.name} ---\n" + "\n".join(rows))
        return _truncate("\n\n".join(sheets), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract XLS {}", path)
        return f"[error: failed to extract XLS: {e!s}]"
    finally:
        if workbook is not None:
            workbook.release_resources()


def _format_xls_cell(cell: Any, workbook: Any, xlrd: Any) -> str:
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return ""
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, f"#ERROR({cell.value})")
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
        return value.isoformat(sep=" ")
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        number = float(cell.value)
        return str(int(number)) if number.is_integer() else str(number)
    return str(cell.value)


def _extract_xlsx(path: Path) -> str:
    """Extract text from XLSX using openpyxl."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[error: openpyxl not installed]"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets: list[str] = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows: list[str] = []
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    sheets.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
            return _truncate("\n\n".join(sheets), _MAX_TEXT_LENGTH)
        finally:
            wb.close()
    except Exception as e:
        logger.exception("Failed to extract XLSX {}", path)
        return f"[error: failed to extract XLSX: {e!s}]"


def _extract_pptx(path: Path) -> str:
    """Extract text from PPTX using python-pptx."""
    try:
        from pptx import Presentation as PptxPresentation
    except ImportError:
        return "[error: python-pptx not installed]"
    try:
        prs = PptxPresentation(path)
        slides: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                _collect_pptx_shape_text(shape, slide_text)
            if slide_text:
                slides.append(f"--- Slide {i} ---\n" + "\n".join(slide_text))
        return _truncate("\n\n".join(slides), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract PPTX {}", path)
        return f"[error: failed to extract PPTX: {e!s}]"


def _collect_pptx_shape_text(shape, out: list[str]) -> None:
    """Collect text from a PPTX shape, recursing into groups and tables.

    Groups have ``has_text_frame=False`` and must be walked via ``.shapes``;
    tables are GraphicFrame objects whose cell text lives under ``.table``.
    """
    sub_shapes = getattr(shape, "shapes", None)
    if sub_shapes is not None:
        for sub in sub_shapes:
            _collect_pptx_shape_text(sub, out)
        return

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                out.append(line)
        return

    text = getattr(shape, "text", "")
    if text:
        out.append(text)


def _extract_text_file(path: Path) -> str:
    """Extract text from a plain text file."""
    try:
        # Try UTF-8 first, then latin-1 fallback
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1")
        return _truncate(content, _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to read text file {}", path)
        return f"[error: failed to read file: {e!s}]"


def _truncate(text: str, max_length: int) -> str:
    """Truncate text with a suffix indicating truncation."""
    if len(text) <= max_length:
        return text
    return text[:max_length] + f"... (truncated, {len(text)} chars total)"


def _is_text_extension(ext: str) -> bool:
    """Check if extension is a text format."""
    return ext in {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".log",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
    }


# ---------------------------------------------------------------------------
# High-level helper: split media into images + extracted document text
# ---------------------------------------------------------------------------

_MAX_EXTRACT_FILE_SIZE = 50 * 1024 * 1024  # 50 MB


def is_image_file(path: str) -> bool:
    """Check whether *path* looks like an image file.

    Uses magic-byte detection (reads first 16 bytes) with a ``mimetypes``
    extension-based fallback.
    """
    p = Path(path)
    mime: str | None = None
    if p.is_file():
        try:
            with p.open("rb") as f:
                mime = detect_image_mime(f.read(16))
        except OSError:
            mime = None
    if not mime:
        mime = mimetypes.guess_type(path)[0]
    return bool(mime and mime.startswith("image/"))


def reference_non_image_attachments(
    content: str, media: list[str],
) -> tuple[str, list[str]]:
    """Separate images from non-image attachments without reading file content.

    Image paths are preserved for downstream vision-block construction.
    Non-image paths are appended as ``[Attachment: path]`` references.
    """
    image_paths: list[str] = []
    attachment_refs: list[str] = []
    for path in media:
        if is_image_file(path):
            image_paths.append(path)
        else:
            attachment_refs.append(f"[Attachment: {path}]")
    if attachment_refs:
        suffix = "\n".join(attachment_refs)
        content = f"{content}\n\n{suffix}" if content else suffix
    return content, image_paths


def extract_documents(
    text: str,
    media_paths: list[str],
    *,
    max_file_size: int = _MAX_EXTRACT_FILE_SIZE,
) -> tuple[str, list[str]]:
    """Separate images from documents in *media_paths*.

    Documents (PDF, DOCX, XLSX, PPTX, plain-text, …) have their text
    extracted and appended to *text*.  Only image paths are kept in the
    returned list so that downstream layers only need to handle vision
    blocks.

    Files larger than *max_file_size* bytes are skipped with a warning
    to avoid unbounded memory / CPU usage.
    """
    image_paths: list[str] = []
    doc_texts: list[str] = []

    for path_str in media_paths:
        p = Path(path_str)
        if not p.is_file():
            continue

        try:
            size = p.stat().st_size
        except OSError:
            continue
        if size > max_file_size:
            logger.warning(
                "Skipping oversized file for extraction: {} ({:.1f} MB > {} MB limit)",
                p.name, size / (1024 * 1024), max_file_size // (1024 * 1024),
            )
            continue

        if is_image_file(path_str):
            image_paths.append(path_str)
        else:
            extracted = extract_text(p)
            if extracted and not extracted.startswith("[error:"):
                doc_texts.append(f"[File: {p.name}]\n{extracted}")

    if doc_texts:
        text = text + "\n\n" + "\n\n".join(doc_texts)

    return text, image_paths
